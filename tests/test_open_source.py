from __future__ import annotations

import io
import re
import xml.etree.ElementTree as ET
from pathlib import Path

import pytest
from PIL import Image
from openpyxl import load_workbook

from server.app import create_app
from server.localization import PHRASES, localize_html
from server.snipe import Resolution, SnipeClient


class SettingsSnipeStub:
    def __init__(self):
        self.last_scan = ""

    @staticmethod
    def user_display_name(user):
        return str(user.get("name") or user.get("email") or user.get("id"))

    def list_status_labels(self):
        return [
            {"id": 1, "name": "Deployed", "type": "deployable"},
            {"id": 2, "name": "Ready to Deploy", "type": "deployable"},
            {"id": 3, "name": "Pending", "type": "pending"},
            {"id": 4, "name": "Wydane", "type": "deployable"},
            {"id": 5, "name": "Gotowe", "type": "deployable"},
            {"id": 6, "name": "Serwis", "type": "pending"},
        ]

    def resolve_asset(self, scan_value):
        self.last_scan = scan_value
        asset = {
            "id": 64,
            "asset_tag": "TEST-00064",
            "serial": "SERIAL-64",
            "name": "Test asset",
            "model": {"name": "Test model"},
            "manufacturer": {"name": "Test manufacturer"},
            "status_label": {"name": "Deployed", "status_meta": "Ready"},
        }
        return Resolution("matched", [asset], {64: ["QR Snipe-IT"]})

    @staticmethod
    def checkout_block_reason(_asset):
        return None

    @staticmethod
    def get_status_label_exact(name):
        return {"id": 1, "name": name}


@pytest.fixture()
def app(tmp_path):
    return create_app(
        {
            "TESTING": True,
            "SECRET_KEY": "test-secret",
            "ADMIN_USERNAME": "root",
            "ADMIN_PASSWORD": "correct-horse-battery-staple",
            "DATABASE_PATH": str(tmp_path / "bridge.db"),
            "LOGO_PATH": str(tmp_path / "branding" / "logo"),
            "SNIPEIT_BASE_URL": "https://snipe.example.org",
            "SNIPEIT_API_TOKEN": "test-token",
            "APP_LANGUAGE": "en",
            "CSRF_ENABLED": True,
            "DATABASE_BACKUP_INTERVAL_HOURS": 0,
            "SESSION_COOKIE_SECURE": False,
        },
        snipe_client=SettingsSnipeStub(),
    )


@pytest.fixture()
def client(app):
    return app.test_client()


def csrf(body: str) -> str:
    match = re.search(r'name="csrf_token" value="([^"]+)"', body)
    assert match
    return match.group(1)


def admin_login(client):
    token = csrf(client.get("/login").get_data(as_text=True))
    return client.post(
        "/login",
        data={
            "csrf_token": token,
            "username": "root",
            "password": "correct-horse-battery-staple",
        },
        follow_redirects=True,
    )


def test_health_and_fresh_terminal(client):
    assert client.get("/healthz").json["version"] == "1.0.0-rc14"
    terminal = client.get("/terminal", follow_redirects=True)
    body = terminal.get_data(as_text=True)
    assert terminal.status_code == 200
    assert "SCAN PAIRING QR" in body
    assert "SNIPE BRIDGE" in body


def test_bootstrap_admin_and_settings_access(client):
    assert client.get("/admin/settings").status_code == 302
    page = admin_login(client)
    assert page.status_code == 200
    settings = client.get("/admin/settings")
    assert settings.status_code == 200
    body = settings.get_data(as_text=True)
    assert "Application settings" in body
    assert "info-tip" in body
    assert "favicon.png" in body
    assert "Accent" not in body
    assert 'name="snipeit_status_checkout"' in body
    assert "Microsoft Entra ID" in body
    assert "/auth/oidc/callback" in body


def test_invalid_admin_password_is_rejected(client):
    token = csrf(client.get("/login").get_data(as_text=True))
    response = client.post(
        "/login",
        data={"csrf_token": token, "username": "root", "password": "wrong"},
        follow_redirects=True,
    )
    assert "Invalid administrator username or password" in response.get_data(as_text=True)


def test_settings_persist_and_reconfigure_branding(client, app):
    admin_login(client)
    token = csrf(client.get("/admin/settings").get_data(as_text=True))
    image = Image.new("RGB", (32, 16), "#123456")
    payload = io.BytesIO()
    image.save(payload, "PNG")
    payload.seek(0)
    response = client.post(
        "/admin/settings",
        data={
            "csrf_token": token,
            "app_name": "Asset Desk",
            "terminal_title": "ASSET DESK",
            "organization_name": "Example Foundation",
            "app_language": "pl",
            "primary_color": "#112233",
            "header_color": "#101010",
            "snipeit_base_url": "https://assets.example.org/",
            "snipeit_api_token": "new-token",
            "snipeit_custom_field": "_snipeit_inventory_1",
            "snipeit_status_checkout": "Wydane",
            "snipeit_status_ready": "Gotowe",
            "snipeit_status_service": "Serwis",
            "identity_provider_type": "google",
            "identity_provider_name": "Google Workspace",
            "oidc_discovery_url": "https://accounts.google.com/.well-known/openid-configuration",
            "oidc_client_id": "client-id",
            "oidc_client_secret": "client-secret",
            "oidc_redirect_uri": "https://bridge.example.org/auth/oidc/callback",
            "identity_allowed_domains": "example.org",
            "oidc_scopes": "openid profile email",
            "oidc_email_claim": "email",
            "oidc_groups_claim": "groups",
            "oidc_admin_groups": "bridge-admins",
            "logo": (payload, "logo.png"),
        },
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert app.config["APP_NAME"] == "Asset Desk"
    assert app.config["SNIPEIT_BASE_URL"] == "https://assets.example.org"
    assert app.extensions["store"].get_settings()["terminal_title"] == "ASSET DESK"
    polish_body = response.get_data(as_text=True)
    assert "Ustawienia aplikacji" in polish_body
    assert "Zapisz ustawienia" in polish_body
    assert "Application settings" not in polish_body
    assert "operator-toast-success" in polish_body
    assert "}, 4000);" in polish_body
    logo = client.get("/branding/logo")
    assert logo.status_code == 200
    assert logo.mimetype == "image/png"


def test_settings_never_expose_api_token_and_support_preserve_replace_clear(client, app):
    admin_login(client)
    page = client.get("/admin/settings").get_data(as_text=True)
    assert "test-token" not in page
    assert "•" * 12 in page

    def settings_payload(token_value: str, clear: bool = False):
        body = client.get("/admin/settings").get_data(as_text=True)
        data = {
            "csrf_token": csrf(body),
            "app_name": "Snipe Bridge",
            "terminal_title": "SNIPE BRIDGE",
            "organization_name": "Example",
            "app_language": "en",
            "primary_color": "#FFCD05",
            "header_color": "#202124",
            "snipeit_base_url": "https://snipe.example.org",
            "snipeit_api_token": token_value,
            "snipeit_custom_field": "",
            "snipeit_status_checkout": "Deployed",
            "snipeit_status_ready": "Ready to Deploy",
            "snipeit_status_service": "Pending",
            "identity_provider_type": "generic",
            "identity_provider_name": "Company SSO",
            "oidc_discovery_url": "https://id.example.org/.well-known/openid-configuration",
            "oidc_client_id": "client-id",
            "oidc_client_secret": "",
            "oidc_redirect_uri": "https://bridge.example.org/auth/oidc/callback",
            "identity_allowed_domains": "example.org",
            "oidc_scopes": "openid profile email",
            "oidc_email_claim": "email",
            "oidc_groups_claim": "groups",
            "oidc_admin_groups": "bridge-admins",
        }
        if clear:
            data["clear_snipeit_api_token"] = "1"
        return data

    client.post("/admin/settings", data=settings_payload("•" * 12))
    assert app.extensions["store"].get_settings()["snipeit_api_token"] == "test-token"
    client.post("/admin/settings", data=settings_payload("replacement-secret"))
    page = client.get("/admin/settings").get_data(as_text=True)
    assert "replacement-secret" not in page
    assert app.extensions["store"].get_settings()["snipeit_api_token"] == "replacement-secret"
    client.post("/admin/settings", data=settings_payload("•" * 12, clear=True))
    assert app.extensions["store"].get_settings()["snipeit_api_token"] == ""


def test_saved_settings_are_refreshed_on_every_worker_request(tmp_path):
    database = tmp_path / "shared.db"
    config = {
        "TESTING": True,
        "SECRET_KEY": "test",
        "ADMIN_PASSWORD": "test",
        "DATABASE_PATH": str(database),
        "DATABASE_BACKUP_INTERVAL_HOURS": 0,
    }
    first = create_app(config)
    second = create_app(config)
    first.extensions["store"].set_settings(
        {"primary_color": "#663399", "app_language": "pl"}
    )
    second.test_client().get("/healthz")
    assert second.config["PRIMARY_COLOR"] == "#663399"
    assert second.config["APP_LANGUAGE"] == "pl"


def test_non_admin_google_session_cannot_open_settings(client):
    with client.session_transaction() as session:
        session["operator_authenticated"] = True
        session["operator_name"] = "Operator"
        session["operator_email"] = "operator@example.org"
        session["operator_auth_method"] = "google"
    assert client.get("/admin/settings").status_code == 403


def test_settings_never_expose_oidc_secret_and_preserve_mask(client, app):
    app.config["OIDC_CLIENT_SECRET"] = "configured-client-secret"
    app.extensions["store"].set_settings(
        {"oidc_client_secret": "configured-client-secret"}
    )
    admin_login(client)
    page = client.get("/admin/settings").get_data(as_text=True)
    assert "configured-client-secret" not in page
    assert "•" * 12 in page


def test_oidc_admin_group_grants_settings_access(client, app):
    app.config["OIDC_ADMIN_GROUPS"] = "bridge-admins"
    with client.session_transaction() as session:
        session["operator_authenticated"] = True
        session["operator_name"] = "OIDC Admin"
        session["operator_email"] = "oidc@example.org"
        session["operator_auth_method"] = "oidc"
        session["identity_groups"] = ["bridge-admins"]
    assert client.get("/admin/settings").status_code == 200


def test_english_operator_and_terminal_do_not_mix_reported_phrases(client):
    operator = admin_login(client).get_data(as_text=True)
    assert "Scan the displayed QR code with the terminal." in operator
    assert "The code is single-use. It expires in" in operator
    assert "Connect the terminal first." in operator
    assert "Connect terminal, aby" not in operator
    terminal = client.get("/terminal", follow_redirects=True).get_data(as_text=True)
    assert "SCAN PAIRING QR" in terminal
    assert "ZESKANUJ QR PAROWANIA" not in terminal


def test_polish_localization_preserves_asset_tag_and_complete_list_words():
    polish = localize_html(
        "Asset Tag | UAM Asset Tag | REMOVE FROM LIST? | REMOVE FROM LIST",
        "pl",
    )
    assert polish == "Asset Tag | UAM Asset Tag | USUNĄĆ Z LISTY? | USUŃ Z LISTY"
    assert "Sprzęt Tag" not in polish
    assert "LISTAY" not in polish


def test_terminal_keeps_scan_focus_supports_mc31xx_keys_and_hides_scan_hint(client):
    body = client.get("/terminal", follow_redirects=True).get_data(as_text=True)
    assert 'onblur="schedulePrimaryFocus()"' in body
    assert "keyCode == 114 || keyCode == 125" in body
    assert "keyCode == 115 || keyCode == 126" in body
    assert "document.onkeydown = handleTerminalKey;" in body
    assert "document.onkeypress = handleTerminalKey;" not in body
    assert "Asset Tag / Serial / UAM Asset Tag / QR" not in body


def test_printable_r_and_s_cannot_trigger_mc31xx_actions_during_scan(client):
    body = client.get("/terminal", follow_redirects=True).get_data(as_text=True)
    ready = body.split("function terminalReady() {", 1)[1].split("</script>", 1)[0]
    assert "document.onkeydown = handleTerminalKey;" in ready
    assert "document.onkeypress" not in ready
    # IE Mobile reports lowercase r/s as character codes 114/115 on keypress.
    # The global hardware handler must never receive that event type.
    assert '"r" is 114' in ready
    assert '"s" is 115' in ready


def test_terminal_blocks_native_enter_submit_until_scan_value_is_stable(client):
    body = client.get("/terminal", follow_redirects=True).get_data(as_text=True)
    assert "var scanSubmitAuthorized = false;" in body
    assert "scanSubmitTimer = window.setTimeout(checkScanStability, 250);" in body
    assert "if (scanStablePasses >= 4)" in body
    assert "if (!scanSubmitAuthorized)" in body
    assert 'onclick="return authorizeScanButton()"' in body


def test_terminal_does_not_submit_during_a_pause_between_scan_characters(client):
    body = client.get("/terminal", follow_redirects=True).get_data(as_text=True)
    handler = body.split("function handleScanEnter(eventObject) {", 1)[1].split(
        "function makeStatusRequest()", 1
    )[0]
    enter_branch, ordinary_character_branch = handler.split(
        "return stopKey(eventValue);", 1
    )
    assert "scheduleScanSubmit();" in enter_branch
    assert "if (scanSubmitTimer !== null)" in ordinary_character_branch
    assert ordinary_character_branch.count("scheduleScanSubmit();") == 0


def test_terminal_shortens_hardware_qr_before_legacy_form_submission(client):
    body = client.get("/terminal", follow_redirects=True).get_data(as_text=True)
    assert "function shortenHardwareQrBeforeSubmit()" in body
    assert "marker = lower.indexOf('/hardware/');" in body
    assert "item.value = '/hardware/' + assetId;" in body
    assert body.count("shortenHardwareQrBeforeSubmit();") >= 3


def test_terminal_status_reflects_mode_change_and_forces_fresh_response(client, app):
    store = app.extensions["store"]
    terminal = store.create_terminal()
    terminal_id = terminal["terminal_id"]
    access_token = terminal.pop("_access_token")
    cookie_name = "stb_terminal_" + terminal_id.replace("-", "")
    client.set_cookie(cookie_name, access_token, domain="localhost")
    paired = store.pair_terminal_setup(
        terminal["pair_code"], "Operator", 1, "operator@example.org"
    )
    assert paired is not None

    first = client.get(f"/terminal/status?t={terminal_id}")
    assert first.status_code == 200
    assert first.get_data(as_text=True).startswith("paired:setup:0:0:-:")
    assert first.headers["Cache-Control"].startswith("no-store")
    assert first.headers["Pragma"] == "no-cache"

    changed = store.reconfigure_terminal_return(
        terminal_id, "ready", "Operator", 1, "operator@example.org", False
    )
    assert changed is not None
    second = client.get(f"/terminal/status?t={terminal_id}")
    assert second.get_data(as_text=True).startswith("paired:checkin:0:0:ready:")
    terminal_page = client.get(f"/terminal?t={terminal_id}").get_data(as_text=True)
    assert "SCAN RETURNED ASSET" in terminal_page


def test_terminal_activity_heartbeat_does_not_change_screen_signature(client, app):
    store = app.extensions["store"]
    terminal = store.create_terminal()
    terminal_id = terminal["terminal_id"]
    access_token = terminal.pop("_access_token")
    client.set_cookie(
        "stb_terminal_" + terminal_id.replace("-", ""),
        access_token,
        domain="localhost",
    )
    assert store.pair_terminal_setup(
        terminal["pair_code"], "Operator", 1, "operator@example.org"
    )
    before = client.get(f"/terminal/status?t={terminal_id}").get_data(as_text=True)
    before_record = store.get_terminal(terminal_id)
    response = client.post(f"/terminal/touch?t={terminal_id}")
    after = client.get(f"/terminal/status?t={terminal_id}").get_data(as_text=True)
    after_record = store.get_terminal(terminal_id)
    assert response.status_code == 204
    assert after == before
    assert after_record["updated_at"] == before_record["updated_at"]
    assert after_record["last_activity_at"] >= before_record["last_activity_at"]


def test_terminal_document_disables_cache_and_mode_reload_uses_unique_url(client):
    response = client.get("/terminal", follow_redirects=True)
    body = response.get_data(as_text=True)
    assert response.headers["Cache-Control"].startswith("no-store")
    assert response.headers["Pragma"] == "no-cache"
    assert "function reloadTerminalDocument()" in body
    assert "?t={{ terminal_id }}" not in body
    assert "&_=' + new Date().getTime()" in body


def test_full_terminal_scan_resolves_hardware_url_segment(client, app):
    class QrResponse:
        status_code = 200

        @staticmethod
        def json():
            return {
                "id": 64,
                "asset_tag": "IT-0064",
                "serial": "SN-0064",
                "model": {"name": "Test model"},
                "status_label": {"name": "Ready", "status_meta": "Ready"},
                "assigned_to": None,
            }

    class QrSession:
        def __init__(self):
            self.urls = []

        def request(self, _method, url, **_kwargs):
            self.urls.append(url)
            assert url.endswith("/api/v1/hardware/64")
            return QrResponse()

    qr_client = SnipeClient("https://configured.example", "token", "")
    qr_session = QrSession()
    qr_client.session = qr_session
    app.extensions["snipe"].resolve_asset = qr_client.resolve_asset

    store = app.extensions["store"]
    terminal = store.create_terminal()
    terminal_id = terminal["terminal_id"]
    access_token = terminal.pop("_access_token")
    assert store.pair_terminal_setup(
        terminal["pair_code"], "Operator", 1, "operator@example.org", "en"
    )
    assert store.reconfigure_terminal_checkout(
        terminal_id,
        77,
        "Recipient",
        "recipient@example.org",
        "Operator",
        1,
        "operator@example.org",
        False,
    )
    client.set_cookie(
        "stb_terminal_" + terminal_id.replace("-", ""),
        access_token,
        domain="localhost",
    )
    scanned_value = (
        "prefix https://another-snipe.example/tenant/hardware/64/details?from=qr suffix"
    )
    response = client.post(
        f"/terminal/scan?t={terminal_id}&field=code",
        data={"code": scanned_value},
    )
    assert response.status_code == 302
    assert qr_session.urls == ["https://configured.example/api/v1/hardware/64"]
    assert store.get_pending(terminal_id)["asset"]["id"] == 64
    confirmation = client.get(response.headers["Location"]).get_data(as_text=True)
    assert "CONFIRM ASSET" in confirmation
    assert "IT-0064" in confirmation


def test_operator_pairing_status_unlocks_operation_selection(client, app):
    admin_login(client)
    client.get("/")
    with client.session_transaction() as session:
        token = session["pending_pair_token"]
    terminal = app.extensions["store"].create_terminal()
    claimed = app.extensions["store"].claim_pair_request(
        token, terminal["terminal_id"]
    )
    assert claimed is not None
    status = client.get("/operator/pairing-qr/status")
    assert status.json["status"] == "paired"
    panel = client.get("/").get_data(as_text=True)
    assert "Choose operation" in panel


def test_pairing_qr_image_is_valid_svg_and_expires_after_claim(client, app):
    admin_login(client)
    with client.session_transaction() as session:
        token = session["pending_pair_token"]
    response = client.get("/operator/pairing-qr/image", query_string={"token": token})
    assert response.status_code == 200
    assert response.mimetype == "image/svg+xml"
    root = ET.fromstring(response.data)
    assert root.tag == "{http://www.w3.org/2000/svg}svg"
    assert root.findall(".//{http://www.w3.org/2000/svg}rect")
    assert response.headers["Cache-Control"] == "no-store"
    terminal = app.extensions["store"].create_terminal()
    assert app.extensions["store"].claim_pair_request(token, terminal["terminal_id"])
    expired = client.get("/operator/pairing-qr/image", query_string={"token": token})
    assert expired.status_code in (404, 410)


@pytest.mark.parametrize("source_format", ["PNG", "JPEG"])
def test_terminal_image_is_scaled_jpeg_with_white_margins(client, app, monkeypatch, source_format):
    source = Image.new("RGBA" if source_format == "PNG" else "RGB", (160, 80), "red")
    if source_format == "PNG":
        source.putpixel((0, 0), (0, 0, 0, 0))
    payload = io.BytesIO()
    source.save(payload, source_format)
    snipe = app.extensions["snipe"]
    monkeypatch.setattr(snipe, "asset_image_url", lambda asset: "https://example.invalid/image", raising=False)
    monkeypatch.setattr(snipe, "fetch_asset_image", lambda asset: payload.getvalue(), raising=False)
    store = app.extensions["store"]
    terminal = store.create_terminal()
    terminal_id = terminal["terminal_id"]
    client.set_cookie("stb_terminal_" + terminal_id.replace("-", ""), terminal["_access_token"])
    store.set_pending(terminal_id, {"id": 64}, "TEST-64", ["Asset Tag"], None)
    response = client.get("/terminal/image", query_string={"t": terminal_id, "asset": 64})
    assert response.status_code == 200
    assert response.mimetype == "image/jpeg"
    with Image.open(io.BytesIO(response.data)) as result:
        result.load()
        assert result.format == "JPEG"
        assert result.size == (78, 72)
        assert result.mode == "RGB"
        assert all(channel >= 240 for channel in result.getpixel((39, 2)))
        red, green, blue = result.getpixel((39, 36))
        assert red > 200 and green < 30 and blue < 30


def test_dynamic_server_messages_are_localized_as_complete_phrases():
    english = localize_html(
        "Brak statusu „Wydane” w Snipe-IT. Lokalnie: Awaryjne logowanie administratora PIN-em.",
        "en",
    )
    assert "Missing status „Wydane” in Snipe-IT." in english
    assert "Legacy administrator PIN sign-in." in english
    assert (
        localize_html("Terminal przełączony do wydania dla: Jan Kowalski.", "en")
        == "Terminal switched to checkout for: Jan Kowalski."
    )
    assert "No asset found for code:" not in localize_html(
        "Nie znaleziono sprzętu dla zeskanowanego kodu.", "en"
    )
    assert (
        localize_html(
            "Sprzęt nie jest przypisany. Nie można go przyjąć.", "en"
        )
        == "The asset is not assigned and cannot be returned."
    )
    assert (
        localize_html("Zakończono masowe wydanie:", "en")
        == "Completed batch checkout:"
    )
    assert (
        localize_html("Zakończono masowy zwrot:", "en")
        == "Completed batch return:"
    )


def test_localization_handles_quantities_and_text_after_protected_tags():
    assert localize_html("ZATWIERDŹ LISTĘ (1 SZT.)", "en") == (
        "CONFIRM LIST (1 ITEM)"
    )
    assert localize_html("ZATWIERDŹ LISTĘ (3 SZT.)", "en") == (
        "CONFIRM LIST (3 ITEMS)"
    )
    assert localize_html("1 szt.", "en") == "1 item"
    assert localize_html("8 szt.", "en") == "8 items"
    assert localize_html("CONFIRM LIST (1 ITEM)", "pl") == (
        "ZATWIERDŹ LISTĘ (1 SZT.)"
    )
    assert localize_html("CONFIRM LIST (8 ITEMS)", "pl") == (
        "ZATWIERDŹ LISTĘ (8 SZT.)"
    )
    assert localize_html(
        '<option value="ready">Gotowy do wydania</option>', "en"
    ) == '<option value="ready">Ready to deploy</option>'


def test_all_template_ui_literals_have_an_english_translation():
    polish_ui = re.compile(
        r"[ąćęłńóśźżĄĆĘŁŃÓŚŹŻ]"
        r"|\b(?:aby|albo|anuluj|brak|jest|kod|krok|listy|może|nie|odbiorca|"
        r"oczekiwanie|operacji|parowania|połącz|przejdź|sprzęt|terminala|"
        r"urządzeń|wpisów|wybierz|wygaśnie|zapisz|zeskanuj|został|zwrotu|szt)\b",
        re.IGNORECASE,
    )
    templates = Path(__file__).parents[1] / "server" / "templates"
    for path in sorted(templates.glob("*.html")):
        source = path.read_text(encoding="utf-8")
        for expression in (
            "{{ total }}",
            "{{ batch_items|length }}",
            "{{ message_count }}",
        ):
            source = source.replace(expression, "1")
        translated = localize_html(source, "en")
        translated = re.sub(r"{[#%].*?[#%]}|{{.*?}}", " ", translated, flags=re.S)
        translated = re.sub(r"/\*.*?\*/", " ", translated, flags=re.S)
        assert not polish_ui.search(translated), (
            f"Untranslated Polish UI text remains in {path.name}"
        )


def test_english_batch_terminal_uses_english_singular_item(client, app):
    store = app.extensions["store"]
    terminal = store.create_terminal()
    terminal_id = terminal["terminal_id"]
    access_token = terminal.pop("_access_token")
    assert store.pair_terminal_setup(
        terminal["pair_code"], "Operator", 1, "operator@example.org", "en"
    )
    assert store.reconfigure_terminal_checkout(
        terminal_id,
        77,
        "Recipient",
        "recipient@example.org",
        "Operator",
        1,
        "operator@example.org",
        True,
    )
    assert store.add_batch_item(
        terminal_id,
        {
            "asset": {
                "id": 64,
                "asset_tag": "TEST-00064",
                "serial": "SERIAL-64",
                "name": "Test asset",
                "model": {"name": "Test model"},
            },
            "scan_value": "TEST-00064",
            "match_fields": "Asset Tag",
        },
    )
    client.set_cookie(
        "stb_terminal_" + terminal_id.replace("-", ""),
        access_token,
        domain="localhost",
    )
    body = client.get(f"/terminal?t={terminal_id}").get_data(as_text=True)
    assert "CONFIRM LIST (1 ITEM)" in body
    assert "1 item" in body
    assert "szt." not in body.lower()
    assert "SZT." not in body


def test_translation_catalog_is_unique_and_round_trips_every_phrase():
    polish = [source for source, _target in PHRASES]
    english = [target for _source, target in PHRASES]
    assert len(polish) == len(set(polish))
    assert len(english) == len(set(english))
    for source, target in PHRASES:
        assert localize_html(source, "en") == target
        assert localize_html(target, "pl") == source


def test_reported_terminal_phrases_are_complete_in_both_languages():
    assert localize_html("PARUJ", "en") == "PAIR"
    assert (
        localize_html("Kod QR wygasł albo został już użyty.", "en")
        == "The QR code has expired or has already been used."
    )
    assert localize_html("Serial Number", "pl") == "Numerze Seryjnym"
    assert localize_html("TRYB ZWROTU", "en") == "RETURN MODE"
    assert "Sprzęt Tag" not in localize_html("Asset Tag", "pl")


def test_language_choice_is_kept_for_operator_and_isolated_between_clients(app):
    polish_client = app.test_client()
    polish_client.get("/language/pl?next=/login")
    assert "Zaloguj się" in polish_client.get("/login").get_data(as_text=True)
    admin_login(polish_client)
    assert "Historia operacji" in polish_client.get("/").get_data(as_text=True)

    english_client = app.test_client()
    assert "Sign in" in english_client.get("/login").get_data(as_text=True)
    assert "Zaloguj się" not in english_client.get("/login").get_data(as_text=True)


def test_qr_pairing_propagates_operator_language_to_terminal(client, app):
    store = app.extensions["store"]
    terminal = store.create_terminal()
    terminal_id = terminal["terminal_id"]
    access_token = terminal.pop("_access_token")
    token = store.create_pair_request(
        "setup", False, "Operator", 1, "operator@example.org", language="pl"
    )
    cookie_name = "stb_terminal_" + terminal_id.replace("-", "")
    client.set_cookie(cookie_name, access_token, domain="localhost")
    response = client.post(
        f"/terminal/pair-qr?t={terminal_id}&field=pairing_code",
        data={"pairing_code": f"STBPAIR1:{token}"},
    )
    assert response.status_code == 302
    assert store.get_terminal(terminal_id)["language"] == "pl"
    body = client.get(f"/terminal?t={terminal_id}").get_data(as_text=True)
    assert "TERMINAL SPAROWANY" in body
    assert "TERMINAL CONNECTED" not in body


def test_long_asset_qr_reaches_snipe_it_resolver_without_truncation(client, app):
    store = app.extensions["store"]
    terminal = store.create_terminal()
    terminal_id = terminal["terminal_id"]
    access_token = terminal.pop("_access_token")
    assert store.pair_terminal_setup(
        terminal["pair_code"], "Operator", 1, "operator@example.org", "en"
    )
    assert store.reconfigure_terminal_checkout(
        terminal_id,
        77,
        "Recipient",
        "recipient@example.org",
        "Operator",
        1,
        "operator@example.org",
        False,
    )
    cookie_name = "stb_terminal_" + terminal_id.replace("-", "")
    client.set_cookie(cookie_name, access_token, domain="localhost")
    scanned_url = "https://snipe.example.org/hardware/64?scanner_payload=" + "a" * 900
    response = client.post(
        f"/terminal/scan?t={terminal_id}&field=code",
        data={"code": scanned_url},
    )
    assert response.status_code == 302
    assert app.extensions["snipe"].last_scan == scanned_url
    assert store.get_pending(terminal_id)["scan_value"] == scanned_url


def test_status_refresh_uses_four_second_popup(client):
    admin_login(client)
    settings = client.get("/admin/settings").get_data(as_text=True)
    response = client.post(
        "/admin/settings/refresh-statuses",
        data={"csrf_token": csrf(settings)},
        follow_redirects=True,
    )
    body = response.get_data(as_text=True)
    assert "Refreshed the Snipe-IT status list: 6." in body
    assert "operator-toast-success" in body
    assert "}, 4000);" in body


def test_login_description_and_identity_warning_can_be_disabled(client, app):
    app.config["LOGIN_DESCRIPTION_ENABLED"] = "0"
    app.config["SHOW_IDENTITY_CONFIG_WARNING"] = "0"
    body = client.get("/login").get_data(as_text=True)
    assert "speed up checkouts and returns" not in body
    assert "Identity-provider sign-in requires" not in body


def test_snipe_status_option_values_are_not_translated(client):
    admin_login(client)
    client.get("/language/pl?next=/admin/settings")
    body = client.get("/admin/settings").get_data(as_text=True)
    assert 'value="Ready to Deploy"' in body
    assert "Ustawienia aplikacji" in body


def test_operator_and_terminal_views_have_reciprocal_screen_guards(client):
    admin_login(client)
    operator = client.get("/").get_data(as_text=True)
    assert 'id="operator-screen-guard"' in operator
    assert 'href="/terminal"' in operator
    terminal = client.get("/terminal", follow_redirects=True).get_data(as_text=True)
    assert 'id="terminal-screen-guard"' in terminal
    assert "OPEN OPERATOR PANEL" in terminal


def test_operator_page_on_windows_ce_renders_only_fixed_terminal_guard(client):
    body = client.get(
        "/login",
        headers={"User-Agent": "Mozilla/4.0 (compatible; MSIE 6.0; Windows CE; MC319Z)"},
    ).get_data(as_text=True)
    assert '<body class="operator-guard-active">' in body
    assert 'id="operator-screen-guard"' in body
    assert 'class="login-grid"' not in body
    assert 'name="username"' not in body
    assert 'name="password"' not in body


def test_unpaired_terminal_language_switch_is_inside_header(client):
    body = client.get("/terminal", follow_redirects=True).get_data(as_text=True)
    header_start = body.index('class="header-table"')
    header_end = body.index("</table>", header_start)
    switch = body.index('class="terminal-language-switch"')
    pairing_panel = body.index('class="panel center pairing-panel"')
    assert header_start < switch < header_end < pairing_panel


def test_settings_show_only_description_for_current_interface_language(client):
    admin_login(client)
    english = client.get("/admin/settings").get_data(as_text=True)
    assert "Sign-in page description" in english
    assert "Opis na stronie logowania" not in english
    assert 'type="hidden" name="login_description_pl"' in english

    client.get("/language/pl?next=/admin/settings")
    polish = client.get("/admin/settings").get_data(as_text=True)
    assert "Opis na stronie logowania" in polish
    assert "Sign-in page description" not in polish
    assert 'type="hidden" name="login_description_en"' in polish


def test_local_admin_audit_message_matches_username_password_login(client):
    admin_login(client)
    body = client.get("/operator/logs").get_data(as_text=True)
    assert "Local administrator sign-in with username and password." in body
    assert "Legacy administrator PIN sign-in." not in body


def test_operation_exports_follow_the_operator_language(client):
    admin_login(client)
    csv_response = client.get("/operator/logs/export/csv")
    csv_body = csv_response.get_data().decode("utf-8-sig")
    assert "operation-log-" in csv_response.headers["Content-Disposition"]
    assert csv_body.startswith("Time;Result;Operation;Asset Tag;Scanned code;")
    assert "Local administrator sign-in with username and password." in csv_body
    assert "Czas;Wynik;Operacja" not in csv_body

    xlsx_response = client.get("/operator/logs/export/xlsx")
    workbook = load_workbook(io.BytesIO(xlsx_response.get_data()), read_only=True)
    assert workbook.active.title == "Operation log"
    assert workbook.active.cell(1, 1).value == "Time"
    assert workbook.active.cell(1, 11).value == "Message"

    client.get("/language/pl?next=/operator/logs")
    polish_response = client.get("/operator/logs/export/csv")
    polish_body = polish_response.get_data().decode("utf-8-sig")
    assert "dziennik-operacji-" in polish_response.headers["Content-Disposition"]
    assert polish_body.startswith("Czas;Wynik;Operacja;Asset Tag;Kod skanu;")
    assert "Lokalne logowanie administratora loginem i hasłem." in polish_body


def test_rc1_organization_specific_custom_field_is_cleared(tmp_path):
    database = tmp_path / "legacy-rc1.db"
    first = create_app(
        {
            "TESTING": True,
            "SECRET_KEY": "test",
            "ADMIN_PASSWORD": "test",
            "DATABASE_PATH": str(database),
            "DATABASE_BACKUP_INTERVAL_HOURS": 0,
        }
    )
    first.extensions["store"].set_settings(
        {"snipeit_custom_field": "_snipeit_uam_asset_tag_7"}
    )
    second = create_app(
        {
            "TESTING": True,
            "SECRET_KEY": "test",
            "ADMIN_PASSWORD": "test",
            "DATABASE_PATH": str(database),
            "DATABASE_BACKUP_INTERVAL_HOURS": 0,
        }
    )
    assert second.config["SNIPEIT_CUSTOM_FIELD"] == ""
    assert second.extensions["store"].get_settings()["snipeit_custom_field"] == ""
