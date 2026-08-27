from __future__ import annotations

import csv
import hmac
import math
import os
from pathlib import Path
import re
import secrets
import sqlite3
import time
from datetime import datetime
from functools import wraps
from io import BytesIO, StringIO
from typing import Any, Callable
from urllib.parse import urlencode

from authlib.integrations.base_client.errors import OAuthError
from authlib.integrations.flask_client import OAuth
from flask import (
    Flask,
    Response,
    flash,
    g,
    jsonify,
    make_response,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from PIL import Image, ImageOps, UnidentifiedImageError
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.graphics import renderSVG
from reportlab.graphics.barcode.qr import QrCodeWidget
from reportlab.graphics.shapes import Drawing
from werkzeug.middleware.proxy_fix import ProxyFix

from .snipe import SnipeClient, SnipeError
from .storage import Store
from .localization import localize_html


BUILD_VERSION = "1.0.0-rc14"
TERMINAL_ID_PATTERN = re.compile(
    r"^[0-9a-f]{8}-[0-9a-f]{4}-[1-5][0-9a-f]{3}-[89ab][0-9a-f]{3}-[0-9a-f]{12}$",
    re.IGNORECASE,
)
PAIR_QR_PREFIX = "STBPAIR1:"
PAIR_TOKEN_PATTERN = re.compile(r"^[A-Za-z0-9_-]{20,64}$")
API_TOKEN_MASK = "•" * 12
AUDIT_PER_PAGE_OPTIONS = (25, 50, 100, 200)
AUDIT_RESULT_LABELS = {
    "success": "Sukces",
    "duplicate": "Duplikat",
    "warning": "Ostrzeżenie",
    "blocked": "Zablokowano",
    "ambiguous": "Niejednoznaczny",
    "not_found": "Nie znaleziono",
    "error": "Błąd",
}
AUDIT_ACTION_LABELS = {
    "login": "Logowanie",
    "pair_qr": "Parowanie QR",
    "manual_pair": "Parowanie ręczne",
    "terminal_replace": "Zmiana terminala",
    "pair_cancel": "Anulowanie parowania",
    "mode_change": "Zmiana trybu",
    "operator_logout": "Wylogowanie operatora",
    "terminal_unpair": "Zakończenie sesji terminala",
    "resolve": "Wyszukanie sprzętu",
    "batch_add": "Dodanie do listy",
    "batch_remove": "Usunięcie z listy",
    "batch_clear": "Wyczyszczenie listy",
    "operation_lock": "Blokada równoległej operacji",
    "diagnostic": "Diagnostyka terminala",
    "checkout": "Wydanie",
    "checkin": "Zwrot",
}

TRANSLATIONS = {
    "en": {
        "history": "Operation history",
        "settings": "Settings",
        "logout": "Sign out",
        "operator_panel": "Operator panel",
        "login_title": "Sign in to start an asset operation",
        "google_login": "Sign in with Google Workspace",
        "admin_login": "Administrator sign-in",
        "username": "Username",
        "password": "Password",
        "sign_in": "Sign in",
    },
    "pl": {
        "history": "Historia operacji",
        "settings": "Ustawienia",
        "logout": "Wyloguj",
        "operator_panel": "Panel operatora",
        "login_title": "Zaloguj się, aby rozpocząć operację na sprzęcie",
        "google_login": "Zaloguj przez Google Workspace",
        "admin_login": "Logowanie administratora",
        "username": "Nazwa użytkownika",
        "password": "Hasło",
        "sign_in": "Zaloguj",
    },
}


def env_bool(name: str, default: bool = True) -> bool:
    value = os.getenv(name)
    if value is None:
        return default
    return value.strip().lower() in {"1", "true", "yes", "on"}


def parse_list(value: str) -> set[str]:
    return {
        item.strip().lstrip("@").casefold()
        for item in re.split(r"[,;\s]+", value)
        if item.strip()
    }


def validate_oidc_identity(
    userinfo: dict[str, Any], email_claim: str, allowed_domains: str
) -> str:
    claim = email_claim.strip() or "email"
    email = str(userinfo.get(claim) or "").strip().casefold()
    if not email or "@" not in email:
        raise ValueError("Dostawca tożsamości nie zwrócił prawidłowego adresu e-mail.")
    if userinfo.get("email_verified") is False:
        raise ValueError("Dostawca tożsamości nie potwierdził adresu e-mail.")
    domains = parse_list(allowed_domains)
    domain = email.rsplit("@", 1)[1]
    if domains and domain not in domains:
        allowed = ", ".join(f"@{item}" for item in sorted(domains))
        raise ValueError(f"Dostęp jest ograniczony do domen: {allowed}.")
    return email


def parse_admin_emails(value: str) -> set[str]:
    return parse_list(value)


def related_name(value: Any) -> str:
    if isinstance(value, dict):
        return str(value.get("name") or "").strip()
    return str(value or "").strip()


def asset_model_label(asset: dict[str, Any]) -> str:
    model = asset.get("model")
    model_name = related_name(model)
    manufacturer_name = related_name(asset.get("manufacturer"))
    if not manufacturer_name and isinstance(model, dict):
        manufacturer_name = related_name(model.get("manufacturer"))

    if manufacturer_name and model_name:
        manufacturer_key = manufacturer_name.casefold()
        model_key = model_name.casefold()
        if model_key == manufacturer_key or model_key.startswith(
            f"{manufacturer_key} "
        ):
            return model_name
        return f"{manufacturer_name} {model_name}"
    return model_name or manufacturer_name or "Brak modelu"


def create_app(
    config_override: dict[str, Any] | None = None,
    *,
    snipe_client: SnipeClient | None = None,
) -> Flask:
    app = Flask(__name__)
    app.config.update(
        SECRET_KEY=os.getenv("SECRET_KEY", "development-only-change-me"),
        ADMIN_PIN=os.getenv("ADMIN_PIN", ""),
        ADMIN_USERNAME=os.getenv("ADMIN_USERNAME", "admin"),
        ADMIN_PASSWORD=os.getenv("ADMIN_PASSWORD", ""),
        ADMIN_EMAILS=os.getenv("ADMIN_EMAILS", ""),
        DATABASE_PATH=os.getenv("DATABASE_PATH", "data/bridge.db"),
        SNIPEIT_BASE_URL=os.getenv("SNIPEIT_BASE_URL", ""),
        SNIPEIT_API_TOKEN=os.getenv("SNIPEIT_API_TOKEN", ""),
        SNIPEIT_CUSTOM_FIELD=os.getenv("SNIPEIT_CUSTOM_FIELD", ""),
        SNIPEIT_STATUS_CHECKOUT=os.getenv("SNIPEIT_STATUS_CHECKOUT", "Deployed"),
        SNIPEIT_STATUS_READY=os.getenv(
            "SNIPEIT_STATUS_READY", "Ready to Deploy"
        ),
        SNIPEIT_STATUS_SERVICE=os.getenv("SNIPEIT_STATUS_SERVICE", "Pending"),
        GOOGLE_CLIENT_ID=os.getenv("GOOGLE_CLIENT_ID", ""),
        GOOGLE_CLIENT_SECRET=os.getenv("GOOGLE_CLIENT_SECRET", ""),
        GOOGLE_ALLOWED_DOMAIN=os.getenv("GOOGLE_ALLOWED_DOMAIN", ""),
        GOOGLE_REDIRECT_URI=os.getenv("GOOGLE_REDIRECT_URI", ""),
        IDENTITY_PROVIDER_TYPE=os.getenv(
            "IDENTITY_PROVIDER_TYPE",
            "google" if os.getenv("GOOGLE_CLIENT_ID") else "generic",
        ),
        IDENTITY_PROVIDER_NAME=os.getenv(
            "IDENTITY_PROVIDER_NAME",
            "Google Workspace" if os.getenv("GOOGLE_CLIENT_ID") else "Identity Provider",
        ),
        OIDC_DISCOVERY_URL=os.getenv(
            "OIDC_DISCOVERY_URL",
            (
                "https://accounts.google.com/.well-known/openid-configuration"
                if os.getenv("GOOGLE_CLIENT_ID")
                else ""
            ),
        ),
        OIDC_CLIENT_ID=os.getenv(
            "OIDC_CLIENT_ID", os.getenv("GOOGLE_CLIENT_ID", "")
        ),
        OIDC_CLIENT_SECRET=os.getenv(
            "OIDC_CLIENT_SECRET", os.getenv("GOOGLE_CLIENT_SECRET", "")
        ),
        OIDC_REDIRECT_URI=os.getenv(
            "OIDC_REDIRECT_URI", os.getenv("GOOGLE_REDIRECT_URI", "")
        ),
        IDENTITY_ALLOWED_DOMAINS=os.getenv(
            "IDENTITY_ALLOWED_DOMAINS", os.getenv("GOOGLE_ALLOWED_DOMAIN", "")
        ),
        OIDC_SCOPES=os.getenv("OIDC_SCOPES", "openid profile email"),
        OIDC_EMAIL_CLAIM=os.getenv("OIDC_EMAIL_CLAIM", "email"),
        OIDC_GROUPS_CLAIM=os.getenv("OIDC_GROUPS_CLAIM", "groups"),
        OIDC_ADMIN_GROUPS=os.getenv("OIDC_ADMIN_GROUPS", ""),
        # The displayed build must describe the code inside the image. Keeping
        # it independent from stack variables prevents an old cached image
        # from presenting itself as a newer release.
        BRIDGE_VERSION=BUILD_VERSION,
        DEVICE_LABEL=os.getenv("DEVICE_LABEL", "Scanner-01"),
        APP_NAME=os.getenv("APP_NAME", "Snipe Bridge"),
        TERMINAL_TITLE=os.getenv("TERMINAL_TITLE", "SNIPE BRIDGE"),
        ORGANIZATION_NAME=os.getenv("ORGANIZATION_NAME", "Your organization"),
        APP_LANGUAGE=os.getenv("APP_LANGUAGE", "en"),
        LOGIN_DESCRIPTION_ENABLED=os.getenv("LOGIN_DESCRIPTION_ENABLED", "1"),
        LOGIN_DESCRIPTION_PL=os.getenv(
            "LOGIN_DESCRIPTION_PL",
            "Podłącz ręczny skaner kodów do Snipe-IT i przyśpiesz wydania oraz zwroty",
        ),
        LOGIN_DESCRIPTION_EN=os.getenv(
            "LOGIN_DESCRIPTION_EN",
            "Connect a handheld barcode scanner to Snipe-IT and speed up checkouts and returns",
        ),
        SHOW_IDENTITY_CONFIG_WARNING=os.getenv(
            "SHOW_IDENTITY_CONFIG_WARNING", "1"
        ),
        PRIMARY_COLOR=os.getenv("PRIMARY_COLOR", "#FFCD05"),
        ACCENT_COLOR="#198754",
        HEADER_COLOR=os.getenv("HEADER_COLOR", "#202124"),
        LOGO_PATH=os.getenv("LOGO_PATH", "/data/branding/logo"),
        ASSET_IMAGES_ENABLED=env_bool("ASSET_IMAGES_ENABLED", True),
        TERMINAL_SAFE_MODE=env_bool("TERMINAL_SAFE_MODE", False),
        TERMINAL_COOKIE_SECURE=env_bool("TERMINAL_COOKIE_SECURE", False),
        PAIR_TTL_MINUTES=int(os.getenv("PAIR_TTL_MINUTES", "15")),
        TERMINAL_IDLE_MINUTES=int(os.getenv("TERMINAL_IDLE_MINUTES", "15")),
        QR_PAIR_TTL_MINUTES=int(os.getenv("QR_PAIR_TTL_MINUTES", "2")),
        REQUEST_TIMEOUT_SECONDS=float(os.getenv("REQUEST_TIMEOUT_SECONDS", "12")),
        STATUS_CACHE_SECONDS=int(os.getenv("STATUS_CACHE_SECONDS", "300")),
        VERIFY_TLS=env_bool("VERIFY_TLS", True),
        TRUST_PROXY=env_bool("TRUST_PROXY", False),
        CSRF_ENABLED=env_bool("CSRF_ENABLED", True),
        LOGIN_MAX_ATTEMPTS=int(os.getenv("LOGIN_MAX_ATTEMPTS", "5")),
        LOGIN_WINDOW_SECONDS=int(os.getenv("LOGIN_WINDOW_SECONDS", "300")),
        LOGIN_BLOCK_SECONDS=int(os.getenv("LOGIN_BLOCK_SECONDS", "900")),
        PAIR_MAX_ATTEMPTS=int(os.getenv("PAIR_MAX_ATTEMPTS", "10")),
        PAIR_WINDOW_SECONDS=int(os.getenv("PAIR_WINDOW_SECONDS", "300")),
        PAIR_BLOCK_SECONDS=int(os.getenv("PAIR_BLOCK_SECONDS", "600")),
        SQLITE_BUSY_TIMEOUT_SECONDS=int(
            os.getenv("SQLITE_BUSY_TIMEOUT_SECONDS", "10")
        ),
        SQLITE_WAL_ENABLED=env_bool("SQLITE_WAL_ENABLED", True),
        AUDIT_RETENTION_DAYS=int(os.getenv("AUDIT_RETENTION_DAYS", "365")),
        TERMINAL_RETENTION_DAYS=int(os.getenv("TERMINAL_RETENTION_DAYS", "90")),
        DATABASE_BACKUP_DIR=os.getenv("DATABASE_BACKUP_DIR", ""),
        DATABASE_BACKUP_INTERVAL_HOURS=int(
            os.getenv("DATABASE_BACKUP_INTERVAL_HOURS", "24")
        ),
        DATABASE_BACKUP_KEEP=int(os.getenv("DATABASE_BACKUP_KEEP", "7")),
        TERMINAL_OPERATION_LOCK_SECONDS=int(
            os.getenv("TERMINAL_OPERATION_LOCK_SECONDS", "120")
        ),
        READINESS_CHECK_SNIPEIT=env_bool("READINESS_CHECK_SNIPEIT", False),
        REQUIRE_SECURE_CONFIG=env_bool("REQUIRE_SECURE_CONFIG", False),
        SESSION_COOKIE_SECURE=env_bool("SESSION_COOKIE_SECURE", False),
        SESSION_COOKIE_HTTPONLY=True,
        SESSION_COOKIE_SAMESITE="Lax",
    )
    if config_override:
        app.config.update(config_override)
    if app.config["TRUST_PROXY"]:
        app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)
    app.jinja_env.globals["asset_model_label"] = asset_model_label

    store = Store(
        app.config["DATABASE_PATH"],
        app.config["PAIR_TTL_MINUTES"],
        app.config["TERMINAL_IDLE_MINUTES"],
        app.config["QR_PAIR_TTL_MINUTES"],
        app.config["SQLITE_BUSY_TIMEOUT_SECONDS"],
        app.config["SQLITE_WAL_ENABLED"],
        app.config["AUDIT_RETENTION_DAYS"],
        app.config["TERMINAL_RETENTION_DAYS"],
        app.config["DATABASE_BACKUP_DIR"],
        app.config["DATABASE_BACKUP_INTERVAL_HOURS"],
        app.config["DATABASE_BACKUP_KEEP"],
        app.config["TERMINAL_OPERATION_LOCK_SECONDS"],
    )

    # Environment variables bootstrap a fresh installation. Values saved by an
    # administrator intentionally take precedence on subsequent starts.
    managed_settings = {
        "APP_NAME": "app_name",
        "TERMINAL_TITLE": "terminal_title",
        "ORGANIZATION_NAME": "organization_name",
        "APP_LANGUAGE": "app_language",
        "LOGIN_DESCRIPTION_ENABLED": "login_description_enabled",
        "LOGIN_DESCRIPTION_PL": "login_description_pl",
        "LOGIN_DESCRIPTION_EN": "login_description_en",
        "SHOW_IDENTITY_CONFIG_WARNING": "show_identity_config_warning",
        "PRIMARY_COLOR": "primary_color",
        "HEADER_COLOR": "header_color",
        "SNIPEIT_BASE_URL": "snipeit_base_url",
        "SNIPEIT_API_TOKEN": "snipeit_api_token",
        "SNIPEIT_CUSTOM_FIELD": "snipeit_custom_field",
        "SNIPEIT_STATUS_CHECKOUT": "snipeit_status_checkout",
        "SNIPEIT_STATUS_READY": "snipeit_status_ready",
        "SNIPEIT_STATUS_SERVICE": "snipeit_status_service",
        "IDENTITY_PROVIDER_TYPE": "identity_provider_type",
        "IDENTITY_PROVIDER_NAME": "identity_provider_name",
        "OIDC_DISCOVERY_URL": "oidc_discovery_url",
        "OIDC_CLIENT_ID": "oidc_client_id",
        "OIDC_CLIENT_SECRET": "oidc_client_secret",
        "OIDC_REDIRECT_URI": "oidc_redirect_uri",
        "IDENTITY_ALLOWED_DOMAINS": "identity_allowed_domains",
        "OIDC_SCOPES": "oidc_scopes",
        "OIDC_EMAIL_CLAIM": "oidc_email_claim",
        "OIDC_GROUPS_CLAIM": "oidc_groups_claim",
        "OIDC_ADMIN_GROUPS": "oidc_admin_groups",
    }
    saved_settings = store.get_settings()
    if (
        "identity_allowed_domains" not in saved_settings
        and saved_settings.get("google_allowed_domain")
    ):
        saved_settings["identity_allowed_domains"] = saved_settings[
            "google_allowed_domain"
        ]
        store.set_settings(
            {"identity_allowed_domains": saved_settings["identity_allowed_domains"]}
        )
    for config_name, setting_name in managed_settings.items():
        if setting_name in saved_settings:
            app.config[config_name] = saved_settings[setting_name]
    if app.config["SNIPEIT_CUSTOM_FIELD"] == "_snipeit_uam_asset_tag_7":
        app.config["SNIPEIT_CUSTOM_FIELD"] = ""
        store.set_settings({"snipeit_custom_field": ""})

    if app.config["REQUIRE_SECURE_CONFIG"]:
        configuration_errors = []
        if app.config["SECRET_KEY"] == "development-only-change-me":
            configuration_errors.append("SECRET_KEY")
        if not str(app.config["ADMIN_PASSWORD"] or app.config["ADMIN_PIN"]).strip():
            configuration_errors.append("ADMIN_PASSWORD")
        if configuration_errors:
            raise RuntimeError(
                "Missing secure configuration: " + ", ".join(configuration_errors)
            )
    snipe = snipe_client or SnipeClient(
        app.config["SNIPEIT_BASE_URL"],
        app.config["SNIPEIT_API_TOKEN"],
        app.config["SNIPEIT_CUSTOM_FIELD"],
        timeout=app.config["REQUEST_TIMEOUT_SECONDS"],
        verify_tls=app.config["VERIFY_TLS"],
        status_cache_seconds=app.config["STATUS_CACHE_SECONDS"],
    )
    app.extensions["store"] = store
    app.extensions["snipe"] = snipe
    app.extensions["managed_settings"] = managed_settings

    def refresh_runtime_settings() -> None:
        """Synchronize settings changed by another Gunicorn worker."""
        nonlocal snipe
        latest = store.get_settings()
        old_connection = (
            app.config["SNIPEIT_BASE_URL"],
            app.config["SNIPEIT_API_TOKEN"],
            app.config["SNIPEIT_CUSTOM_FIELD"],
        )
        for config_name, setting_name in managed_settings.items():
            if setting_name in latest:
                app.config[config_name] = latest[setting_name]
        new_connection = (
            app.config["SNIPEIT_BASE_URL"],
            app.config["SNIPEIT_API_TOKEN"],
            app.config["SNIPEIT_CUSTOM_FIELD"],
        )
        if new_connection != old_connection and snipe_client is None:
            snipe = SnipeClient(
                app.config["SNIPEIT_BASE_URL"],
                app.config["SNIPEIT_API_TOKEN"],
                app.config["SNIPEIT_CUSTOM_FIELD"],
                timeout=app.config["REQUEST_TIMEOUT_SECONDS"],
                verify_tls=app.config["VERIFY_TLS"],
                status_cache_seconds=app.config["STATUS_CACHE_SECONDS"],
            )
            app.extensions["snipe"] = snipe

    def normalized_language(value: Any) -> str:
        language = str(value or "").strip().lower()
        return language if language in {"pl", "en"} else ""

    def setting_enabled(value: Any) -> bool:
        return str(value or "").strip().lower() in {"1", "true", "yes", "on"}

    def is_legacy_terminal_request() -> bool:
        user_agent = str(request.headers.get("User-Agent") or "").casefold()
        return any(
            marker in user_agent
            for marker in ("windows ce", "mc319", "mc31")
        )

    def request_language() -> str:
        forced = normalized_language(getattr(g, "response_language", ""))
        if forced:
            return forced
        selected = normalized_language(session.get("language"))
        if selected:
            return selected
        selected = normalized_language(request.cookies.get("bridge_language"))
        if selected:
            return selected
        return normalized_language(app.config.get("APP_LANGUAGE")) or "en"

    def translate(key: str) -> str:
        language = request_language()
        catalog = TRANSLATIONS.get(language, TRANSLATIONS["en"])
        return catalog.get(key, TRANSLATIONS["en"].get(key, key))

    @app.context_processor
    def inject_application_theme() -> dict[str, Any]:
        logo_path = Path(str(app.config["LOGO_PATH"]))
        return {
            "app_name": app.config["APP_NAME"],
            "terminal_title": app.config["TERMINAL_TITLE"],
            "organization_name": app.config["ORGANIZATION_NAME"],
            "app_language": request_language(),
            "primary_color": app.config["PRIMARY_COLOR"],
            "accent_color": app.config["ACCENT_COLOR"],
            "header_color": app.config["HEADER_COLOR"],
            "custom_logo_available": logo_path.is_file(),
            "tr": translate,
            "operator_is_admin": operator_is_admin(),
            "login_description_enabled": setting_enabled(
                app.config["LOGIN_DESCRIPTION_ENABLED"]
            ),
            "login_description": (
                app.config["LOGIN_DESCRIPTION_PL"]
                if request_language() == "pl"
                else app.config["LOGIN_DESCRIPTION_EN"]
            ),
            "show_identity_config_warning": setting_enabled(
                app.config["SHOW_IDENTITY_CONFIG_WARNING"]
            ),
            "operator_terminal_guard": is_legacy_terminal_request(),
        }

    identity_oauth = OAuth(app)
    identity_client = None
    identity_client_signature: tuple[str, ...] | None = None

    def identity_enabled() -> bool:
        return bool(
            str(app.config["OIDC_DISCOVERY_URL"]).strip()
            and str(app.config["OIDC_CLIENT_ID"]).strip()
            and str(app.config["OIDC_CLIENT_SECRET"]).strip()
        )

    def get_identity_client() -> Any:
        nonlocal identity_client, identity_client_signature
        signature = (
            str(app.config["OIDC_DISCOVERY_URL"]).strip(),
            str(app.config["OIDC_CLIENT_ID"]).strip(),
            str(app.config["OIDC_CLIENT_SECRET"]).strip(),
            str(app.config["OIDC_SCOPES"]).strip(),
        )
        if not identity_enabled():
            identity_client = None
            identity_client_signature = signature
            return None
        if identity_client is None or identity_client_signature != signature:
            identity_oauth._clients.pop("identity", None)
            identity_client = identity_oauth.register(
                "identity",
                overwrite=True,
                client_id=signature[1],
                client_secret=signature[2],
                server_metadata_url=signature[0],
                client_kwargs={"scope": signature[3] or "openid profile email"},
            )
            identity_client_signature = signature
        return identity_client

    app.extensions["identity_oauth"] = identity_oauth
    image_cache: dict[tuple[int, str], bytes] = {}

    def operator_required(view: Callable[..., Any]) -> Callable[..., Any]:
        @wraps(view)
        def wrapped(*args: Any, **kwargs: Any) -> Any:
            if not session.get("operator_authenticated"):
                return redirect(url_for("login", next=request.path))
            return view(*args, **kwargs)

        return wrapped

    def admin_required(view: Callable[..., Any]) -> Callable[..., Any]:
        @wraps(view)
        def wrapped(*args: Any, **kwargs: Any) -> Any:
            if not session.get("operator_authenticated"):
                return redirect(url_for("login", next=request.path))
            if not operator_is_admin():
                return Response("Administrator access required.", status=403)
            return view(*args, **kwargs)

        return wrapped

    def parse_terminal_id() -> str | None:
        terminal_id = (request.values.get("t") or "").strip()
        if not terminal_id or not TERMINAL_ID_PATTERN.match(terminal_id):
            return None
        return terminal_id

    def terminal_cookie_name(terminal_id: str) -> str:
        return "stb_terminal_" + terminal_id.replace("-", "")

    def get_terminal_id() -> str | None:
        terminal_id = parse_terminal_id()
        if not terminal_id:
            return None
        access_token = str(request.cookies.get(terminal_cookie_name(terminal_id)) or "")
        if not store.verify_terminal_access(terminal_id, access_token):
            return None
        return terminal_id

    def set_terminal_cookie(response: Any, terminal_id: str, token: str) -> Any:
        response.set_cookie(
            terminal_cookie_name(terminal_id),
            token,
            max_age=31536000,
            httponly=True,
            secure=bool(app.config["TERMINAL_COOKIE_SECURE"]),
            samesite="Lax",
            path="/terminal",
        )
        return response

    def terminal_redirect(terminal_id: str, **params: str) -> Any:
        query = {"t": terminal_id, **params}
        return redirect(f"{url_for('terminal')}?{urlencode(query)}")

    def terminal_text_value(expected_name: str) -> str:
        value = request.form.get(expected_name)
        if value is None:
            value = next(
                (
                    request.form.get(key, "")
                    for key in request.form.keys()
                    if key.startswith("scan_")
                ),
                "",
            )
        return (
            str(value or "")
            .replace("\r", "")
            .replace("\n", "")
            .strip()[:2048]
        )

    def operator_session_values() -> tuple[str, int | None, str]:
        operator_name = str(session.get("operator_name") or "operator")
        raw_id = session.get("operator_user_id")
        operator_user_id = int(raw_id) if raw_id is not None else None
        operator_email = str(session.get("operator_email") or "")
        return operator_name, operator_user_id, operator_email

    def operator_is_admin() -> bool:
        if session.get("operator_auth_method") == "pin":
            return True
        operator_email = str(session.get("operator_email") or "").casefold()
        if operator_email and operator_email in parse_admin_emails(
            str(app.config["ADMIN_EMAILS"])
        ):
            return True
        session_groups = {
            str(group).strip().casefold()
            for group in session.get("identity_groups", [])
            if str(group).strip()
        }
        return bool(session_groups & parse_list(str(app.config["OIDC_ADMIN_GROUPS"])))

    def csrf_token() -> str:
        token = str(session.get("csrf_token") or "")
        if not token:
            token = secrets.token_urlsafe(32)
            session["csrf_token"] = token
        return token

    app.jinja_env.globals["csrf_token"] = csrf_token

    @app.before_request
    def synchronize_saved_settings() -> None:
        refresh_runtime_settings()

    @app.before_request
    def verify_csrf() -> Any:
        if (
            not app.config["CSRF_ENABLED"]
            or request.method != "POST"
            or request.path.startswith("/terminal/")
        ):
            return None
        supplied = str(
            request.form.get("csrf_token")
            or request.headers.get("X-CSRF-Token")
            or ""
        )
        expected = str(session.get("csrf_token") or "")
        if not expected or not hmac.compare_digest(supplied, expected):
            return Response("Nieprawidłowy token formularza.", status=400)
        return None

    def rate_limit_key(kind: str) -> str:
        return f"{kind}:{request.remote_addr or 'unknown'}"

    def rate_limit_message(retry_after: int) -> str:
        seconds = max(1, int(retry_after))
        return f"Zbyt wiele prób. Spróbuj ponownie za {seconds} s."

    def audit_scope() -> tuple[int | None, str]:
        if operator_is_admin():
            return None, ""
        _, operator_user_id, operator_email = operator_session_values()
        return operator_user_id, operator_email

    def audit_filter_request() -> tuple[dict[str, str], dict[str, Any], list[str]]:
        raw = {
            key: request.args.get(key, "").strip()
            for key in (
                "date_from",
                "date_to",
                "result",
                "action",
                "asset",
                "target",
                "operator",
                "message",
            )
        }
        compiled: dict[str, Any] = {
            key: raw[key]
            for key in ("result", "action", "asset", "target", "operator", "message")
            if raw[key]
        }
        errors: list[str] = []
        for raw_key, compiled_key, end_of_day in (
            ("date_from", "created_from", False),
            ("date_to", "created_to", True),
        ):
            if not raw[raw_key]:
                continue
            try:
                timestamp = int(
                    datetime.strptime(raw[raw_key], "%Y-%m-%d").timestamp()
                )
                compiled[compiled_key] = timestamp + (86400 if end_of_day else 0)
            except ValueError:
                errors.append(f"Nieprawidłowa data: {raw[raw_key]}.")
        return raw, compiled, errors

    def audit_result_tone(result: str) -> str:
        normalized = result.strip().casefold()
        if normalized == "success":
            return "success"
        if normalized in {
            "duplicate",
            "warning",
            "blocked",
            "ambiguous",
            "not_found",
        }:
            return "warning"
        if normalized in {"error", "failed", "failure"}:
            return "error"
        return "neutral"

    def prepare_audit_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
        for row in rows:
            row["created_display"] = datetime.fromtimestamp(row["created_at"]).strftime(
                "%Y-%m-%d %H:%M:%S"
            )
            result = str(row.get("result") or "unknown")
            row["result_label"] = AUDIT_RESULT_LABELS.get(result, result)
            row["result_tone"] = audit_result_tone(result)
            action = str(row.get("action") or "unknown")
            row["action_label"] = AUDIT_ACTION_LABELS.get(action, action)
        return rows

    def audit_export_values(row: dict[str, Any]) -> list[Any]:
        return [
            row.get("created_display") or "",
            row.get("result_label") or row.get("result") or "",
            row.get("action_label") or row.get("action") or "",
            row.get("asset_tag") or "",
            row.get("scan_value") or "",
            row.get("match_fields") or "",
            row.get("target_name") or "",
            row.get("operator_name") or "",
            row.get("operator_email") or "",
            row.get("operator_user_id") or "",
            row.get("message") or "",
            row.get("terminal_id") or "",
        ]

    def localized_audit_export_values(row: dict[str, Any]) -> list[Any]:
        values = audit_export_values(row)
        language = request_language()
        # Result, action, match method and message are application-owned text.
        # Asset, recipient and operator data must remain unchanged.
        for index in (1, 2, 5, 10):
            values[index] = localize_html(str(values[index]), language)
        return values

    def pair_request_belongs_to_operator(pair_request: dict[str, Any]) -> bool:
        operator_name, operator_user_id, operator_email = operator_session_values()
        request_user_id = pair_request.get("operator_user_id")
        request_email = str(pair_request.get("operator_email") or "").casefold()
        request_name = str(pair_request.get("operator_name") or "").casefold()
        if request_user_id is not None and operator_user_id is not None:
            return int(request_user_id) == int(operator_user_id)
        if operator_email and request_email:
            return request_email == operator_email.casefold()
        return bool(operator_name) and request_name == operator_name.casefold()

    def get_session_pair_request() -> tuple[str, dict[str, Any]] | None:
        token = str(session.get("pending_pair_token") or "")
        if not PAIR_TOKEN_PATTERN.match(token):
            session.pop("pending_pair_token", None)
            return None
        pair_request = store.get_pair_request(token)
        if not pair_request or not pair_request_belongs_to_operator(pair_request):
            session.pop("pending_pair_token", None)
            return None
        return token, pair_request

    def terminal_log(terminal_record: dict[str, Any], **event: Any) -> None:
        store.log_event(
            terminal_id=terminal_record.get("terminal_id"),
            operator_name=terminal_record.get("operator_name"),
            operator_user_id=terminal_record.get("operator_user_id"),
            operator_email=terminal_record.get("operator_email"),
            target_user_id=terminal_record.get("target_user_id"),
            target_name=terminal_record.get("target_name"),
            **event,
        )

    def activate_operator_terminal(
        terminal_record: dict[str, Any], *, pairing_action: str
    ) -> None:
        previous_terminal_id = str(session.get("active_terminal_id") or "")
        new_terminal_id = str(terminal_record["terminal_id"])
        if (
            TERMINAL_ID_PATTERN.match(previous_terminal_id)
            and previous_terminal_id != new_terminal_id
        ):
            previous_terminal = store.get_terminal(previous_terminal_id)
            if previous_terminal and previous_terminal["is_paired"]:
                terminal_log(
                    previous_terminal,
                    action="terminal_replace",
                    result="success",
                    message=f"Terminal zastąpiony przez {new_terminal_id}.",
                )
                store.clear_target(previous_terminal_id)
        session["active_terminal_id"] = new_terminal_id
        pending_pair_token = str(session.pop("pending_pair_token", "") or "")
        if PAIR_TOKEN_PATTERN.match(pending_pair_token):
            store.cancel_pair_request(pending_pair_token)
        terminal_log(
            terminal_record,
            action=pairing_action,
            result="success",
            message="Terminal połączony z panelem operatora.",
        )

    def operation_note(terminal_record: dict[str, Any]) -> str:
        parts = [f"Snipe Bridge {app.config['DEVICE_LABEL']}"]
        operator_name = str(terminal_record.get("operator_name") or "").strip()
        operator_email = str(terminal_record.get("operator_email") or "").strip()
        operator_user_id = terminal_record.get("operator_user_id")
        if operator_name:
            parts.append(f"operator: {operator_name}")
        if operator_email:
            parts.append(f"e-mail: {operator_email}")
        if operator_user_id is not None:
            parts.append(f"Snipe ID: {operator_user_id}")
        return "; ".join(parts)

    def execute_terminal_item(
        terminal_record: dict[str, Any], item: dict[str, Any]
    ) -> tuple[str, str]:
        """Revalidate and execute one checkout/checkin prepared by a scan."""
        asset = item["asset"]
        operation = str(asset.get("_bridge_operation") or "checkout")
        current_operation = str(terminal_record["operation_mode"] or "checkout")
        if operation != current_operation:
            raise SnipeError(
                "Tryb terminala zmienił się. Zeskanuj sprzęt ponownie."
            )
        if operation == "checkout" and (
            not terminal_record["target_user_id"]
            or int(terminal_record["target_user_id"])
            != int(item["target_user_id"])
        ):
            raise SnipeError(
                "Odbiorca zmienił się przed potwierdzeniem. Zeskanuj ponownie."
            )

        fresh_resolution = snipe.resolve_asset(str(item["scan_value"]))
        if fresh_resolution.status != "matched":
            raise SnipeError(
                "Nie można ponownie jednoznacznie potwierdzić sprzętu. Zeskanuj go jeszcze raz."
            )
        fresh_asset = fresh_resolution.candidates[0]
        if int(fresh_asset["id"]) != int(asset["id"]):
            raise SnipeError(
                "Kod wskazuje teraz inny sprzęt. Zeskanuj go jeszcze raz."
            )

        note = operation_note(terminal_record)
        asset_label = asset.get("asset_tag") or asset.get("name") or asset["id"]
        if operation == "checkout":
            block_reason = snipe.checkout_block_reason(fresh_asset)
            if block_reason:
                raise SnipeError(block_reason)
            snipe.checkout_asset(
                int(asset["id"]),
                int(terminal_record["target_user_id"]),
                int(asset["_bridge_status_id"]),
                note,
            )
            return "checkout", f"Przypisano {asset_label}."

        block_reason = snipe.checkin_block_reason(
            fresh_asset,
            (
                app.config["SNIPEIT_STATUS_CHECKOUT"],
                app.config["SNIPEIT_STATUS_READY"],
            ),
        )
        if block_reason:
            raise SnipeError(block_reason)
        if asset.get("_bridge_return_choice") == "service":
            note = f"{note}; Wymaga serwisu"
        snipe.checkin_asset(
            int(asset["id"]), int(asset["_bridge_status_id"]), note
        )
        return "checkin", f"Przyjęto {asset_label}."

    @app.after_request
    def no_cache_terminal(response: Any) -> Any:
        if response.mimetype == "text/html" and response.direct_passthrough is False:
            response.set_data(
                localize_html(response.get_data(as_text=True), request_language())
            )
        if (
            request.path.startswith("/terminal")
            and request.path != "/terminal/image"
        ) or request.path in {
            "/operator/pairing-qr/status",
            "/operator/active-terminal/status",
        }:
            response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate"
            response.headers["Pragma"] = "no-cache"
            response.headers["Expires"] = "0"
        return response

    @app.get("/healthz")
    def healthz() -> Any:
        return jsonify(
            {
                "status": "ok",
                "version": app.config["BRIDGE_VERSION"],
                "terminal_safe_mode": bool(app.config["TERMINAL_SAFE_MODE"]),
            }
        )

    @app.get("/readyz")
    def readyz() -> Any:
        checks: dict[str, Any] = {}
        errors: list[str] = []
        try:
            database_check = store.readiness()
            checks["storage"] = database_check
            if database_check.get("database") != "ok":
                errors.append("database")
            if database_check.get("maintenance_error"):
                errors.append("database_backup")
        except sqlite3.Error as exc:
            checks["storage"] = {"database": "error", "message": str(exc)}
            errors.append("database")

        configuration_errors = []
        if app.config["SECRET_KEY"] == "development-only-change-me":
            configuration_errors.append("SECRET_KEY")
        if not str(app.config["SNIPEIT_API_TOKEN"]).strip():
            configuration_errors.append("SNIPEIT_API_TOKEN")
        if not (app.config["ADMIN_PASSWORD"] or app.config["ADMIN_PIN"]) and not identity_enabled():
            configuration_errors.append("operator_authentication")
        checks["configuration"] = (
            {"status": "ok"}
            if not configuration_errors
            else {"status": "error", "missing": configuration_errors}
        )
        if configuration_errors:
            errors.append("configuration")

        if app.config["READINESS_CHECK_SNIPEIT"]:
            try:
                snipe.test_connection()
                checks["snipeit"] = {"status": "ok"}
            except SnipeError as exc:
                checks["snipeit"] = {"status": "error", "message": str(exc)}
                errors.append("snipeit")

        status_code = 503 if errors else 200
        return (
            jsonify(
                {
                    "status": "ready" if not errors else "not_ready",
                    "version": app.config["BRIDGE_VERSION"],
                    "checks": checks,
                }
            ),
            status_code,
        )

    @app.route("/login", methods=["GET", "POST"])
    def login() -> Any:
        if request.method == "POST":
            login_key = rate_limit_key("login")
            retry_after = store.rate_limit_retry_after(login_key)
            if retry_after:
                store.log_event(
                    operator_name=request.form.get("operator_name", "")[:160],
                    action="login",
                    result="blocked",
                    message="Logowanie PIN zablokowane limitem prób.",
                )
                flash(rate_limit_message(retry_after), "error")
                return render_template(
                    "login.html",
                    identity_enabled=identity_enabled(),
                    identity_provider_name=app.config["IDENTITY_PROVIDER_NAME"],
                    identity_allowed_domains=sorted(
                        parse_list(str(app.config["IDENTITY_ALLOWED_DOMAINS"]))
                    ),
                )
            configured_password = str(
                app.config["ADMIN_PASSWORD"] or app.config["ADMIN_PIN"]
            )
            supplied_password = str(
                request.form.get("password") or request.form.get("pin") or ""
            )
            configured_username = str(app.config["ADMIN_USERNAME"] or "admin")
            operator_name = str(
                request.form.get("username")
                or request.form.get("operator_name")
                or ""
            ).strip()
            # ADMIN_PIN remains a compatibility fallback for existing POC test
            # harnesses. New installations should use ADMIN_PASSWORD and the
            # explicit ADMIN_USERNAME.
            valid_username = (
                hmac.compare_digest(operator_name, configured_username)
                if app.config["ADMIN_PASSWORD"]
                else bool(operator_name)
            )
            valid_password = bool(configured_password) and hmac.compare_digest(
                supplied_password, configured_password
            )
            if not configured_password:
                flash("Ustaw ADMIN_PASSWORD w zmiennych kontenera.", "error")
            elif valid_username and valid_password:
                selected_language = request_language()
                store.clear_rate_limit(login_key)
                store.log_event(
                    operator_name=operator_name,
                    operator_email=operator_name if "@" in operator_name else "",
                    action="login",
                    result="success",
                    message="Lokalne logowanie administratora loginem i hasłem.",
                )
                session.clear()
                session["language"] = selected_language
                session["operator_authenticated"] = True
                session["operator_name"] = operator_name
                session["operator_email"] = (
                    operator_name if "@" in operator_name else ""
                )
                session["operator_auth_method"] = "pin"
                return redirect(url_for("operator_panel"))
            else:
                retry_after = store.register_rate_limit_failure(
                    login_key,
                    max_attempts=app.config["LOGIN_MAX_ATTEMPTS"],
                    window_seconds=app.config["LOGIN_WINDOW_SECONDS"],
                    block_seconds=app.config["LOGIN_BLOCK_SECONDS"],
                )
                flash("Nieprawidłowa nazwa administratora lub hasło.", "error")
                store.log_event(
                    operator_name=operator_name,
                    operator_email=operator_name if "@" in operator_name else "",
                    action="login",
                    result="error",
                    message="Nieudane lokalne logowanie administratora.",
                )
                if retry_after:
                    flash(rate_limit_message(retry_after), "error")
        return render_template(
            "login.html",
            identity_enabled=identity_enabled(),
            identity_provider_name=app.config["IDENTITY_PROVIDER_NAME"],
            identity_allowed_domains=sorted(
                parse_list(str(app.config["IDENTITY_ALLOWED_DOMAINS"]))
            ),
        )

    @app.get("/language/<language>")
    def set_language(language: str) -> Any:
        selected = normalized_language(language)
        if not selected:
            return Response("Unsupported language.", status=400)
        target = str(request.args.get("next") or url_for("login"))
        if not target.startswith("/") or target.startswith("//"):
            target = url_for("login")
        session["language"] = selected
        response = redirect(target)
        response.set_cookie(
            "bridge_language",
            selected,
            max_age=31536000,
            httponly=False,
            secure=bool(app.config["SESSION_COOKIE_SECURE"]),
            samesite="Lax",
        )
        return response

    @app.get("/branding/logo")
    def branding_logo() -> Any:
        logo_path = Path(str(app.config["LOGO_PATH"]))
        if logo_path.is_file():
            return send_file(logo_path, mimetype="image/png")
        return redirect(url_for("static", filename="logo-default.png", v="3"))

    @app.route("/admin/settings", methods=["GET", "POST"])
    @admin_required
    def admin_settings() -> Any:
        nonlocal snipe

        def status_options() -> tuple[list[dict[str, Any]], str]:
            try:
                return snipe.list_status_labels(), ""
            except SnipeError as exc:
                return [], str(exc)

        def render_settings(values: dict[str, str]) -> Any:
            display_values = dict(values)
            display_values["snipeit_api_token"] = (
                API_TOKEN_MASK if values.get("snipeit_api_token") else ""
            )
            display_values["oidc_client_secret"] = (
                API_TOKEN_MASK if values.get("oidc_client_secret") else ""
            )
            statuses, status_error = status_options()
            return render_template(
                "admin_settings.html",
                settings=display_values,
                status_options=statuses,
                status_names={str(item["name"]) for item in statuses},
                status_sync_error=status_error,
                default_redirect_uri=url_for("identity_callback", _external=True),
            )

        if request.method == "POST":
            allowed_languages = {"en", "pl"}
            language = str(request.form.get("app_language") or "en").lower()
            if language not in allowed_languages:
                language = "en"
            provider_type = str(
                request.form.get("identity_provider_type") or "generic"
            ).lower()
            if provider_type not in {
                "google",
                "microsoft",
                "keycloak",
                "authentik",
                "okta",
                "generic",
            }:
                provider_type = "generic"
            color_pattern = re.compile(r"^#[0-9a-fA-F]{6}$")
            values = {
                "app_name": str(request.form.get("app_name") or "").strip()[:120],
                "terminal_title": str(request.form.get("terminal_title") or "").strip()[:32],
                "organization_name": str(request.form.get("organization_name") or "").strip()[:120],
                "app_language": language,
                "login_description_enabled": (
                    "1" if request.form.get("login_description_enabled") == "1" else "0"
                ),
                "login_description_pl": str(
                    request.form.get("login_description_pl") or ""
                ).strip()[:500],
                "login_description_en": str(
                    request.form.get("login_description_en") or ""
                ).strip()[:500],
                "show_identity_config_warning": (
                    "1"
                    if request.form.get("show_identity_config_warning") == "1"
                    else "0"
                ),
                "primary_color": str(request.form.get("primary_color") or "#FFCD05"),
                "header_color": str(request.form.get("header_color") or "#202124"),
                "snipeit_base_url": str(request.form.get("snipeit_base_url") or "").strip().rstrip("/"),
                "snipeit_api_token": str(app.config["SNIPEIT_API_TOKEN"]),
                "snipeit_custom_field": str(request.form.get("snipeit_custom_field") or "").strip(),
                "snipeit_status_checkout": str(request.form.get("snipeit_status_checkout") or "").strip(),
                "snipeit_status_ready": str(request.form.get("snipeit_status_ready") or "").strip(),
                "snipeit_status_service": str(request.form.get("snipeit_status_service") or "").strip(),
                "identity_provider_type": provider_type,
                "identity_provider_name": str(request.form.get("identity_provider_name") or "").strip()[:120],
                "oidc_discovery_url": str(request.form.get("oidc_discovery_url") or "").strip(),
                "oidc_client_id": str(request.form.get("oidc_client_id") or "").strip(),
                "oidc_client_secret": str(app.config["OIDC_CLIENT_SECRET"]),
                "oidc_redirect_uri": str(request.form.get("oidc_redirect_uri") or "").strip(),
                "identity_allowed_domains": ",".join(
                    sorted(parse_list(str(request.form.get("identity_allowed_domains") or "")))
                ),
                "oidc_scopes": str(request.form.get("oidc_scopes") or "openid profile email").strip(),
                "oidc_email_claim": str(request.form.get("oidc_email_claim") or "email").strip(),
                "oidc_groups_claim": str(request.form.get("oidc_groups_claim") or "groups").strip(),
                "oidc_admin_groups": ",".join(
                    sorted(parse_list(str(request.form.get("oidc_admin_groups") or "")))
                ),
            }
            submitted_token = str(request.form.get("snipeit_api_token") or "").strip()
            if request.form.get("clear_snipeit_api_token") == "1":
                values["snipeit_api_token"] = ""
            elif submitted_token and submitted_token != API_TOKEN_MASK:
                values["snipeit_api_token"] = submitted_token
            submitted_secret = str(request.form.get("oidc_client_secret") or "").strip()
            if request.form.get("clear_oidc_client_secret") == "1":
                values["oidc_client_secret"] = ""
            elif submitted_secret and submitted_secret != API_TOKEN_MASK:
                values["oidc_client_secret"] = submitted_secret
            for color_key in ("primary_color", "header_color"):
                if not color_pattern.match(values[color_key]):
                    flash(f"Nieprawidłowa wartość koloru: {color_key}.", "error")
                    return render_settings(values)
            if not values["app_name"]:
                values["app_name"] = "Snipe Bridge"
            if not values["terminal_title"]:
                values["terminal_title"] = "SNIPE BRIDGE"
            if not values["organization_name"]:
                values["organization_name"] = "Your organization"
            if not values["identity_provider_name"]:
                values["identity_provider_name"] = "Identity Provider"
            uploaded_logo = request.files.get("logo")
            if uploaded_logo and uploaded_logo.filename:
                logo_bytes = uploaded_logo.read(2 * 1024 * 1024 + 1)
                if len(logo_bytes) > 2 * 1024 * 1024:
                    flash("Logo jest większe niż 2 MB.", "error")
                    return render_settings(values)
                try:
                    image = Image.open(BytesIO(logo_bytes))
                    image.verify()
                    normalized_logo = Image.open(BytesIO(logo_bytes))
                    normalized_logo.load()
                except (UnidentifiedImageError, OSError):
                    flash("Logo musi być poprawnym obrazem PNG, JPEG, GIF lub WebP.", "error")
                    return render_settings(values)
                logo_path = Path(str(app.config["LOGO_PATH"]))
                logo_path.parent.mkdir(parents=True, exist_ok=True)
                if normalized_logo.mode not in {"RGB", "RGBA"}:
                    normalized_logo = normalized_logo.convert("RGBA")
                normalized_logo.save(logo_path, format="PNG", optimize=True)
            store.set_settings(values)
            for config_name, setting_name in managed_settings.items():
                if setting_name in values:
                    app.config[config_name] = values[setting_name]
            session["language"] = language
            if snipe_client is None:
                snipe = SnipeClient(
                    app.config["SNIPEIT_BASE_URL"],
                    app.config["SNIPEIT_API_TOKEN"],
                    app.config["SNIPEIT_CUSTOM_FIELD"],
                    timeout=app.config["REQUEST_TIMEOUT_SECONDS"],
                    verify_tls=app.config["VERIFY_TLS"],
                    status_cache_seconds=app.config["STATUS_CACHE_SECONDS"],
                )
                app.extensions["snipe"] = snipe
            flash("Ustawienia zostały zapisane.", "success")
            return redirect(url_for("admin_settings"))
        values = {
            setting_name: str(app.config[config_name])
            for config_name, setting_name in managed_settings.items()
        }
        return render_settings(values)

    @app.post("/admin/settings/refresh-statuses")
    @admin_required
    def admin_refresh_statuses() -> Any:
        try:
            statuses = snipe.list_status_labels()
            if statuses:
                flash(
                    f"Odświeżono listę statusów Snipe-IT: {len(statuses)}.",
                    "success",
                )
            else:
                flash(
                    "Snipe-IT zwrócił pustą listę statusów. Zachowano bieżące wartości.",
                    "toast-warning",
                )
        except SnipeError as exc:
            flash(f"Nie udało się odświeżyć listy statusów: {exc}", "toast-error")
        return redirect(url_for("admin_settings"))

    @app.get("/auth/oidc")
    def identity_login() -> Any:
        identity = get_identity_client()
        if not identity:
            flash("Logowanie przez dostawcę tożsamości nie jest jeszcze skonfigurowane.", "error")
            return redirect(url_for("login"))
        redirect_uri = app.config["OIDC_REDIRECT_URI"] or url_for(
            "identity_callback", _external=True
        )
        authorize_kwargs: dict[str, Any] = {}
        domains = sorted(parse_list(str(app.config["IDENTITY_ALLOWED_DOMAINS"])))
        if app.config["IDENTITY_PROVIDER_TYPE"] == "google" and len(domains) == 1:
            authorize_kwargs["hd"] = domains[0]
        return identity.authorize_redirect(redirect_uri, **authorize_kwargs)

    @app.get("/auth/oidc/callback")
    def identity_callback() -> Any:
        identity = get_identity_client()
        if not identity:
            return redirect(url_for("login"))
        try:
            selected_language = request_language()
            token = identity.authorize_access_token()
            userinfo = token.get("userinfo") if isinstance(token, dict) else None
            if not isinstance(userinfo, dict):
                userinfo = identity.userinfo(token=token)
            if not isinstance(userinfo, dict):
                raise ValueError("Dostawca tożsamości nie zwrócił danych zalogowanego konta.")
            email = validate_oidc_identity(
                userinfo,
                str(app.config["OIDC_EMAIL_CLAIM"]),
                str(app.config["IDENTITY_ALLOWED_DOMAINS"]),
            )
            snipe_user = snipe.find_user_by_email(email)
            provider_name = str(app.config["IDENTITY_PROVIDER_NAME"])
            store.log_event(
                operator_name=snipe.user_display_name(snipe_user),
                operator_user_id=int(snipe_user["id"]),
                operator_email=email,
                action="login",
                result="success",
                message=f"Logowanie przez dostawcę tożsamości: {provider_name}.",
            )
            session.clear()
            session["language"] = selected_language
            session["operator_authenticated"] = True
            session["operator_name"] = snipe.user_display_name(snipe_user)
            session["operator_email"] = email
            session["operator_user_id"] = int(snipe_user["id"])
            session["operator_auth_method"] = "oidc"
            groups_value = userinfo.get(str(app.config["OIDC_GROUPS_CLAIM"])) or []
            if isinstance(groups_value, str):
                groups_value = re.split(r"[,;]", groups_value)
            session["identity_groups"] = [
                str(group).strip()
                for group in groups_value
                if str(group).strip()
            ] if isinstance(groups_value, (list, tuple, set)) else []
            return redirect(url_for("operator_panel"))
        except (OAuthError, SnipeError, ValueError, KeyError, TypeError) as exc:
            session.clear()
            flash(f"Nie udało się zalogować: {exc}", "error")
            return redirect(url_for("login"))

    @app.get("/auth/google")
    def google_login() -> Any:
        return redirect(url_for("identity_login"))

    @app.get("/auth/google/callback")
    def google_callback() -> Any:
        return redirect(url_for("identity_callback", **request.args))

    @app.post("/logout")
    def logout() -> Any:
        selected_language = request_language()
        active_terminal_id = str(session.get("active_terminal_id") or "")
        if TERMINAL_ID_PATTERN.match(active_terminal_id):
            terminal_record = store.get_terminal(active_terminal_id)
            if terminal_record and terminal_record["is_paired"]:
                terminal_log(
                    terminal_record,
                    action="operator_logout",
                    result="success",
                    message="Operator wylogował panel i zakończył sesję terminala.",
                )
                store.clear_target(active_terminal_id)

        pending_pair_token = str(session.get("pending_pair_token") or "")
        if PAIR_TOKEN_PATTERN.match(pending_pair_token):
            operator_name, operator_user_id, operator_email = operator_session_values()
            store.log_event(
                operator_name=operator_name,
                operator_user_id=operator_user_id,
                operator_email=operator_email,
                action="pair_cancel",
                result="success",
                message="Anulowano oczekujące parowanie podczas wylogowania.",
            )
            store.cancel_pair_request(pending_pair_token)
        session.clear()
        session["language"] = selected_language
        return redirect(url_for("login"))

    @app.get("/")
    @operator_required
    def operator_panel() -> Any:
        query = request.args.get("q", "").strip()
        active_terminal = None
        pairing_token = ""
        pairing_expires_in = 0
        active_terminal_id = str(session.get("active_terminal_id") or "")
        if active_terminal_id:
            candidate = store.get_terminal(active_terminal_id)
            if candidate and candidate["is_paired"]:
                active_terminal = candidate
            else:
                session.pop("active_terminal_id", None)

        now = int(time.time())
        resolved = get_session_pair_request()
        if resolved:
            token, pair_request = resolved
            claimed_terminal_id = pair_request.get("claimed_terminal_id")
            if pair_request.get("claimed_at") and claimed_terminal_id:
                candidate = store.get_terminal(str(claimed_terminal_id))
                session.pop("pending_pair_token", None)
                if candidate and candidate["is_paired"]:
                    active_terminal = candidate
                    session["active_terminal_id"] = candidate["terminal_id"]
            elif (
                not pair_request.get("cancelled_at")
                and int(pair_request["expires_at"]) >= now
            ):
                pairing_token = token
                pairing_expires_in = max(0, int(pair_request["expires_at"]) - now)
            else:
                session.pop("pending_pair_token", None)

        if active_terminal is None and not pairing_token:
            operator_name, operator_user_id, operator_email = operator_session_values()
            pairing_token = store.create_pair_request(
                "setup",
                False,
                operator_name,
                operator_user_id,
                operator_email,
                language=request_language(),
            )
            session["pending_pair_token"] = pairing_token
            pair_request = store.get_pair_request(pairing_token)
            if pair_request:
                pairing_expires_in = max(0, int(pair_request["expires_at"]) - now)

        mode = request.args.get("mode", "").strip()
        batch_mode = request.args.get("batch", "0").strip() == "1"
        users: list[dict[str, Any]] = []
        replacement_pending = active_terminal is not None and bool(pairing_token)
        using_active_terminal = active_terminal is not None and not pairing_token
        if not using_active_terminal or mode not in {"checkout", "checkin"}:
            mode = ""
        if mode == "checkout" and query:
            try:
                users = snipe.search_users(query)
            except SnipeError as exc:
                flash(str(exc), "error")
        return render_template(
            "operator.html",
            operator_name=session.get("operator_name"),
            operator_email=session.get("operator_email"),
            operator_is_admin=operator_is_admin(),
            users=users,
            query=query,
            mode=mode,
            batch_mode=batch_mode,
            active_terminal=active_terminal,
            using_active_terminal=using_active_terminal,
            replacement_pending=replacement_pending,
            pairing_token=pairing_token,
            pairing_expires_in=pairing_expires_in,
            user_display_name=snipe.user_display_name,
            base_url=app.config["SNIPEIT_BASE_URL"],
        )

    @app.post("/operator/pair-code")
    @operator_required
    def operator_pair_code_setup() -> Any:
        pair_limit_key = rate_limit_key("manual-pair")
        retry_after = store.rate_limit_retry_after(pair_limit_key)
        if retry_after:
            flash(rate_limit_message(retry_after), "error")
            return redirect(url_for("operator_panel"))
        pair_code = request.form.get("pair_code", "").strip()
        if not pair_code.isdigit() or len(pair_code) != 6:
            store.register_rate_limit_failure(
                pair_limit_key,
                max_attempts=app.config["PAIR_MAX_ATTEMPTS"],
                window_seconds=app.config["PAIR_WINDOW_SECONDS"],
                block_seconds=app.config["PAIR_BLOCK_SECONDS"],
            )
            flash("Kod parowania musi mieć sześć cyfr.", "error")
            return redirect(url_for("operator_panel"))
        operator_name, operator_user_id, operator_email = operator_session_values()
        terminal_record = store.pair_terminal_setup(
            pair_code,
            operator_name,
            operator_user_id,
            operator_email,
            request_language(),
        )
        if not terminal_record:
            retry_after = store.register_rate_limit_failure(
                pair_limit_key,
                max_attempts=app.config["PAIR_MAX_ATTEMPTS"],
                window_seconds=app.config["PAIR_WINDOW_SECONDS"],
                block_seconds=app.config["PAIR_BLOCK_SECONDS"],
            )
            flash("Kod parowania wygasł albo nie istnieje.", "error")
            if retry_after:
                flash(rate_limit_message(retry_after), "error")
        else:
            store.clear_rate_limit(pair_limit_key)
            activate_operator_terminal(terminal_record, pairing_action="manual_pair")
            pending_pair_token = str(session.pop("pending_pair_token", "") or "")
            if PAIR_TOKEN_PATTERN.match(pending_pair_token):
                store.cancel_pair_request(pending_pair_token)
            flash("Terminal połączony. Wybierz teraz operację.", "success")
        return redirect(url_for("operator_panel"))

    @app.post("/operator/pair")
    @operator_required
    def operator_pair() -> Any:
        pair_limit_key = rate_limit_key("manual-pair")
        retry_after = store.rate_limit_retry_after(pair_limit_key)
        if retry_after:
            flash(rate_limit_message(retry_after), "error")
            return redirect(url_for("operator_panel"))
        pair_code = request.form.get("pair_code", "").strip()
        user_id_raw = request.form.get("user_id", "").strip()
        batch_mode = request.form.get("batch_mode", "0").strip() == "1"
        if not pair_code.isdigit() or len(pair_code) != 6:
            store.register_rate_limit_failure(
                pair_limit_key,
                max_attempts=app.config["PAIR_MAX_ATTEMPTS"],
                window_seconds=app.config["PAIR_WINDOW_SECONDS"],
                block_seconds=app.config["PAIR_BLOCK_SECONDS"],
            )
            flash("Kod parowania musi mieć sześć cyfr.", "error")
            return redirect(
                url_for("operator_panel", pair_code=pair_code, mode="checkout")
            )
        try:
            user_id = int(user_id_raw)
            user = snipe.get_user(user_id)
            name = snipe.user_display_name(user)
            email = str(user.get("email") or user.get("username") or "")
            operator_name, operator_user_id, operator_email = operator_session_values()
            terminal_record = store.pair_terminal(
                pair_code,
                user_id,
                name,
                email,
                operator_name,
                operator_user_id,
                operator_email,
                batch_mode,
                request_language(),
            )
            if not terminal_record:
                retry_after = store.register_rate_limit_failure(
                    pair_limit_key,
                    max_attempts=app.config["PAIR_MAX_ATTEMPTS"],
                    window_seconds=app.config["PAIR_WINDOW_SECONDS"],
                    block_seconds=app.config["PAIR_BLOCK_SECONDS"],
                )
                flash("Kod parowania wygasł albo nie istnieje.", "error")
                if retry_after:
                    flash(rate_limit_message(retry_after), "error")
            else:
                store.clear_rate_limit(pair_limit_key)
                flash(
                    f"Terminal ustawiony do wydania dla: {name}.",
                    "terminal-ready",
                )
                activate_operator_terminal(
                    terminal_record, pairing_action="manual_pair"
                )
                session.pop("pair_code", None)
                return redirect(url_for("operator_panel"))
        except (ValueError, SnipeError) as exc:
            flash(str(exc), "error")
        return redirect(url_for("operator_panel", pair_code=pair_code, mode="checkout"))

    @app.post("/operator/pair-return")
    @operator_required
    def operator_pair_return() -> Any:
        pair_limit_key = rate_limit_key("manual-pair")
        retry_after = store.rate_limit_retry_after(pair_limit_key)
        if retry_after:
            flash(rate_limit_message(retry_after), "error")
            return redirect(url_for("operator_panel"))
        pair_code = request.form.get("pair_code", "").strip()
        return_choice = request.form.get("return_choice", "ready").strip()
        batch_mode = request.form.get("batch_mode", "0").strip() == "1"
        if not pair_code.isdigit() or len(pair_code) != 6:
            store.register_rate_limit_failure(
                pair_limit_key,
                max_attempts=app.config["PAIR_MAX_ATTEMPTS"],
                window_seconds=app.config["PAIR_WINDOW_SECONDS"],
                block_seconds=app.config["PAIR_BLOCK_SECONDS"],
            )
            flash("Kod parowania musi mieć sześć cyfr.", "error")
            return redirect(
                url_for("operator_panel", pair_code=pair_code, mode="checkin")
            )
        try:
            operator_name, operator_user_id, operator_email = operator_session_values()
            terminal_record = store.pair_terminal_return(
                pair_code,
                return_choice,
                operator_name,
                operator_user_id,
                operator_email,
                batch_mode,
                request_language(),
            )
            if not terminal_record:
                retry_after = store.register_rate_limit_failure(
                    pair_limit_key,
                    max_attempts=app.config["PAIR_MAX_ATTEMPTS"],
                    window_seconds=app.config["PAIR_WINDOW_SECONDS"],
                    block_seconds=app.config["PAIR_BLOCK_SECONDS"],
                )
                flash("Kod parowania wygasł albo nie istnieje.", "error")
                if retry_after:
                    flash(rate_limit_message(retry_after), "error")
            else:
                store.clear_rate_limit(pair_limit_key)
                target_status = (
                    app.config["SNIPEIT_STATUS_READY"]
                    if return_choice == "ready"
                    else app.config["SNIPEIT_STATUS_SERVICE"]
                )
                flash(
                    f"Terminal ustawiony do zwrotu. Status docelowy: {target_status}.",
                    "terminal-ready",
                )
                activate_operator_terminal(
                    terminal_record, pairing_action="manual_pair"
                )
                session.pop("pair_code", None)
                return redirect(url_for("operator_panel"))
        except ValueError as exc:
            flash(str(exc), "error")
        return redirect(url_for("operator_panel", pair_code=pair_code, mode="checkin"))

    @app.post("/operator/reconfigure")
    @operator_required
    def operator_reconfigure() -> Any:
        terminal_id = request.form.get("terminal_id", "").strip()
        operation = request.form.get("operation", "").strip()
        batch_mode = request.form.get("batch_mode", "0").strip() == "1"
        if terminal_id != str(session.get("active_terminal_id") or ""):
            flash("Aktywne parowanie terminala wygasło.", "error")
            return redirect(url_for("operator_panel"))
        try:
            operator_name, operator_user_id, operator_email = operator_session_values()
            if operation == "checkout":
                user_id = int(request.form.get("user_id", "").strip())
                user = snipe.get_user(user_id)
                name = snipe.user_display_name(user)
                email = str(user.get("email") or user.get("username") or "")
                terminal_record = store.reconfigure_terminal_checkout(
                    terminal_id,
                    user_id,
                    name,
                    email,
                    operator_name,
                    operator_user_id,
                    operator_email,
                    batch_mode,
                )
                description = "masowego wydania" if batch_mode else "wydania"
                success_message = f"Terminal przełączony do {description} dla: {name}."
            elif operation == "checkin":
                return_choice = request.form.get("return_choice", "ready").strip()
                terminal_record = store.reconfigure_terminal_return(
                    terminal_id,
                    return_choice,
                    operator_name,
                    operator_user_id,
                    operator_email,
                    batch_mode,
                )
                description = "masowego zwrotu" if batch_mode else "zwrotu"
                success_message = f"Terminal przełączony do {description}."
            else:
                raise ValueError("Nieprawidłowa operacja terminala.")
            if not terminal_record:
                current_terminal = store.get_terminal(terminal_id)
                if current_terminal and current_terminal["is_paired"]:
                    flash(
                        "Terminal wykonuje teraz operację. Spróbuj ponownie po jej zakończeniu.",
                        "error",
                    )
                else:
                    session.pop("active_terminal_id", None)
                    flash("Terminal nie jest już sparowany.", "error")
            else:
                terminal_log(
                    terminal_record,
                    action="mode_change",
                    result="success",
                    message=success_message,
                )
                flash(success_message, "terminal-ready")
        except (ValueError, SnipeError) as exc:
            flash(str(exc), "error")
        return redirect(url_for("operator_panel"))

    @app.post("/operator/pairing-qr/create")
    @operator_required
    def operator_create_pair_request() -> Any:
        try:
            operator_name, operator_user_id, operator_email = operator_session_values()
            active_terminal_id = str(session.get("active_terminal_id") or "")
            replace_terminal_id = None
            if TERMINAL_ID_PATTERN.match(active_terminal_id):
                active_terminal = store.get_terminal(active_terminal_id)
                if active_terminal and active_terminal["is_paired"]:
                    replace_terminal_id = active_terminal_id
            token = store.create_pair_request(
                "setup",
                False,
                operator_name,
                operator_user_id,
                operator_email,
                replace_terminal_id=replace_terminal_id,
                language=request_language(),
            )
            previous_token = str(session.get("pending_pair_token") or "")
            if PAIR_TOKEN_PATTERN.match(previous_token):
                store.cancel_pair_request(previous_token)

            session["pending_pair_token"] = token
            return redirect(url_for("operator_panel"))
        except (SnipeError, ValueError, TypeError) as exc:
            flash(str(exc), "error")
            return redirect(url_for("operator_panel"))

    @app.get("/operator/pairing-qr")
    @operator_required
    def operator_pairing_qr() -> Any:
        resolved = get_session_pair_request()
        if not resolved:
            flash("Nie ma aktywnego kodu QR parowania.", "error")
            return redirect(url_for("operator_panel"))
        token, pair_request = resolved
        now = int(time.time())
        if (
            pair_request.get("cancelled_at")
            or pair_request.get("claimed_at")
            or int(pair_request["expires_at"]) < now
        ):
            session.pop("pending_pair_token", None)
            flash("Kod QR parowania wygasł albo został już użyty.", "error")
            return redirect(url_for("operator_panel"))
        return render_template(
            "pairing_qr.html",
            pair_request=pair_request,
            token=token,
            expires_in=max(0, int(pair_request["expires_at"]) - now),
        )

    @app.get("/operator/pairing-qr/image")
    @operator_required
    def operator_pairing_qr_image() -> Any:
        resolved = get_session_pair_request()
        token = request.args.get("token", "").strip()
        if not resolved or not hmac.compare_digest(token, resolved[0]):
            return Response(status=404)
        pair_request = resolved[1]
        if (
            pair_request.get("cancelled_at")
            or pair_request.get("claimed_at")
            or int(pair_request["expires_at"]) < int(time.time())
        ):
            return Response(status=410)

        qr_widget = QrCodeWidget(f"{PAIR_QR_PREFIX}{token}", barLevel="M")
        x1, y1, x2, y2 = qr_widget.getBounds()
        size = 288
        scale = size / max(x2 - x1, y2 - y1)
        drawing = Drawing(
            size,
            size,
            transform=[scale, 0, 0, scale, -x1 * scale, -y1 * scale],
        )
        drawing.add(qr_widget)
        response = Response(renderSVG.drawToString(drawing), mimetype="image/svg+xml")
        response.headers["Cache-Control"] = "no-store"
        return response

    @app.get("/operator/pairing-qr/status")
    @operator_required
    def operator_pairing_qr_status() -> Any:
        resolved = get_session_pair_request()
        if not resolved:
            return jsonify({"status": "missing"}), 404
        token, pair_request = resolved
        if pair_request.get("claimed_at") and pair_request.get("claimed_terminal_id"):
            terminal_id = str(pair_request["claimed_terminal_id"])
            terminal_record = store.get_terminal(terminal_id)
            session.pop("pending_pair_token", None)
            if terminal_record and terminal_record["is_paired"]:
                session["active_terminal_id"] = terminal_id
                return jsonify({"status": "paired", "terminal_id": terminal_id})
            return jsonify({"status": "expired"})
        if pair_request.get("cancelled_at") or int(pair_request["expires_at"]) < int(
            time.time()
        ):
            session.pop("pending_pair_token", None)
            return jsonify({"status": "expired"})
        return jsonify(
            {
                "status": "waiting",
                "expires_in": max(
                    0, int(pair_request["expires_at"]) - int(time.time())
                ),
            }
        )

    @app.get("/operator/active-terminal/status")
    @operator_required
    def operator_active_terminal_status() -> Any:
        terminal_id = str(session.get("active_terminal_id") or "")
        terminal_record = store.get_terminal(terminal_id) if terminal_id else None
        is_paired = bool(terminal_record and terminal_record["is_paired"])
        if not is_paired:
            session.pop("active_terminal_id", None)
        response = jsonify(
            {
                "status": "paired" if is_paired else "ended",
                "terminal_id": terminal_id if is_paired else "",
            }
        )
        response.headers["Cache-Control"] = "no-store"
        return response

    @app.post("/operator/pairing-qr/cancel")
    @operator_required
    def operator_pairing_qr_cancel() -> Any:
        resolved = get_session_pair_request()
        if resolved:
            store.cancel_pair_request(resolved[0])
            operator_name, operator_user_id, operator_email = operator_session_values()
            store.log_event(
                operator_name=operator_name,
                operator_user_id=operator_user_id,
                operator_email=operator_email,
                action="pair_cancel",
                result="success",
                message="Operator anulował oczekujące parowanie QR.",
            )
        session.pop("pending_pair_token", None)
        return redirect(url_for("operator_panel"))

    @app.post("/operator/test")
    @operator_required
    def operator_test() -> Any:
        try:
            snipe.test_connection()
            for status_name in (
                app.config["SNIPEIT_STATUS_CHECKOUT"],
                app.config["SNIPEIT_STATUS_READY"],
                app.config["SNIPEIT_STATUS_SERVICE"],
            ):
                snipe.get_status_label_exact(status_name)
            flash("Połączenie i wymagane statusy Snipe-IT są poprawne.", "success")
        except SnipeError as exc:
            flash(str(exc), "error")
        return redirect(url_for("operator_panel"))

    @app.get("/operator/logs")
    @operator_required
    def operator_logs() -> Any:
        raw_filters, compiled_filters, filter_errors = audit_filter_request()
        try:
            per_page = int(request.args.get("per_page", "25"))
        except ValueError:
            per_page = 25
        if per_page not in AUDIT_PER_PAGE_OPTIONS:
            per_page = 25
        try:
            page = max(1, int(request.args.get("page", "1")))
        except ValueError:
            page = 1
        scope_user_id, scope_email = audit_scope()
        logs, total = store.list_audit_logs(
            limit=per_page,
            offset=(page - 1) * per_page,
            operator_user_id=scope_user_id,
            operator_email=scope_email,
            filters=compiled_filters,
        )
        total_pages = max(1, math.ceil(total / per_page))
        if page > total_pages:
            page = total_pages
            logs, total = store.list_audit_logs(
                limit=per_page,
                offset=(page - 1) * per_page,
                operator_user_id=scope_user_id,
                operator_email=scope_email,
                filters=compiled_filters,
            )
        prepare_audit_rows(logs)
        filter_options = store.audit_filter_options(
            operator_user_id=scope_user_id,
            operator_email=scope_email,
        )
        active_filter_params = {
            key: value for key, value in raw_filters.items() if value
        }
        page_params = {"per_page": per_page, **active_filter_params}
        previous_url = (
            url_for("operator_logs", page=page - 1, **page_params)
            if page > 1
            else ""
        )
        next_url = (
            url_for("operator_logs", page=page + 1, **page_params)
            if page < total_pages
            else ""
        )
        return render_template(
            "logs.html",
            logs=logs,
            operator_name=session.get("operator_name"),
            operator_email=session.get("operator_email"),
            operator_is_admin=operator_is_admin(),
            filters=raw_filters,
            filter_errors=filter_errors,
            filter_options=filter_options,
            result_labels=AUDIT_RESULT_LABELS,
            action_labels=AUDIT_ACTION_LABELS,
            per_page=per_page,
            per_page_options=AUDIT_PER_PAGE_OPTIONS,
            page=page,
            total=total,
            total_pages=total_pages,
            page_start=((page - 1) * per_page + 1) if total else 0,
            page_end=min(page * per_page, total),
            previous_url=previous_url,
            next_url=next_url,
            export_csv_url=url_for(
                "operator_logs_export", file_format="csv", **active_filter_params
            ),
            export_xlsx_url=url_for(
                "operator_logs_export", file_format="xlsx", **active_filter_params
            ),
        )

    @app.get("/operator/logs/export/<file_format>")
    @operator_required
    def operator_logs_export(file_format: str) -> Any:
        if file_format not in {"csv", "xlsx"}:
            return Response(status=404)
        raw_filters, compiled_filters, filter_errors = audit_filter_request()
        if filter_errors:
            for error in filter_errors:
                flash(error, "error")
            return redirect(
                url_for(
                    "operator_logs",
                    **{key: value for key, value in raw_filters.items() if value},
                )
            )
        scope_user_id, scope_email = audit_scope()
        logs, _ = store.list_audit_logs(
            limit=None,
            operator_user_id=scope_user_id,
            operator_email=scope_email,
            filters=compiled_filters,
        )
        prepare_audit_rows(logs)
        headers = [
            "Czas",
            "Wynik",
            "Operacja",
            "Asset Tag",
            "Kod skanu",
            "Dopasowano po",
            "Odbiorca",
            "Operator",
            "E-mail operatora",
            "Snipe ID operatora",
            "Komunikat",
            "ID terminala",
        ]
        language = request_language()
        headers = [localize_html(header, language) for header in headers]
        timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        export_basename = (
            "dziennik-operacji" if language == "pl" else "operation-log"
        )

        def safe_export_value(value: Any) -> Any:
            if not isinstance(value, str):
                return value
            if value.startswith(("=", "+", "-", "@")):
                return "'" + value
            return value

        if file_format == "csv":
            text_output = StringIO(newline="")
            writer = csv.writer(text_output, delimiter=";", lineterminator="\r\n")
            writer.writerow(headers)
            for row in logs:
                writer.writerow(
                    [
                        safe_export_value(value)
                        for value in localized_audit_export_values(row)
                    ]
                )
            output = BytesIO(("\ufeff" + text_output.getvalue()).encode("utf-8"))
            return send_file(
                output,
                mimetype="text/csv; charset=utf-8",
                as_attachment=True,
                download_name=f"{export_basename}-{timestamp}.csv",
            )

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = localize_html("Dziennik operacji", language)
        worksheet.append(headers)
        header_fill = PatternFill("solid", fgColor="FF6600")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center")
        result_fills = {
            "success": PatternFill("solid", fgColor="C6EFCE"),
            "warning": PatternFill("solid", fgColor="FFEB9C"),
            "error": PatternFill("solid", fgColor="FFC7CE"),
            "neutral": PatternFill("solid", fgColor="E7E6E6"),
        }
        for row in logs:
            worksheet.append(
                [
                    safe_export_value(value)
                    for value in localized_audit_export_values(row)
                ]
            )
            worksheet.cell(worksheet.max_row, 2).fill = result_fills[
                str(row["result_tone"])
            ]
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        widths = (20, 18, 18, 18, 24, 24, 24, 24, 30, 20, 48, 38)
        for index, width in enumerate(widths, start=1):
            worksheet.column_dimensions[
                worksheet.cell(1, index).column_letter
            ].width = width
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return send_file(
            output,
            mimetype=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
            as_attachment=True,
            download_name=f"{export_basename}-{timestamp}.xlsx",
        )

    @app.get("/terminal")
    def terminal() -> Any:
        terminal_id = parse_terminal_id()
        terminal_record = store.get_terminal(terminal_id) if terminal_id else None
        if not terminal_record:
            terminal_record = store.create_terminal()
            response = terminal_redirect(terminal_record["terminal_id"])
            return set_terminal_cookie(
                response,
                terminal_record["terminal_id"],
                terminal_record.pop("_access_token"),
            )

        access_token = str(
            request.cookies.get(terminal_cookie_name(terminal_id)) or ""
        )
        provisioned_token = ""
        if not store.verify_terminal_access(terminal_id, access_token):
            provisioned_token = str(store.provision_terminal_access(terminal_id) or "")
            if not provisioned_token:
                return Response("Brak dostępu do tej sesji terminala.", status=403)

        terminal_language = normalized_language(terminal_record.get("language"))
        if terminal_record["is_paired"] and terminal_language:
            g.response_language = terminal_language

        pending = store.get_pending(terminal_id)
        batch_items = store.get_batch_items(terminal_id)
        batch_remove_item = None
        remove_asset_id = request.args.get("remove_asset", "").strip()
        if remove_asset_id.isdigit():
            requested_asset_id = int(remove_asset_id)
            batch_remove_item = next(
                (
                    item
                    for item in batch_items
                    if int((item.get("asset") or {}).get("id") or 0)
                    == requested_asset_id
                ),
                None,
            )
        result = request.args.get("result", "")
        message = request.args.get("msg", "")[:240]
        message_count = request.args.get("count", "").strip()
        if not message_count.isdigit():
            message_count = ""
        terminal_template = (
            "terminal_safe.html"
            if app.config["TERMINAL_SAFE_MODE"]
            else "terminal.html"
        )
        response = make_response(
            render_template(
                terminal_template,
                terminal=terminal_record,
                pending=pending,
                batch_items=batch_items,
                batch_remove_item=batch_remove_item,
                scan_field_name=f"scan_{secrets.token_hex(8)}",
                terminal_id=terminal_id,
                asset_images_enabled=app.config["ASSET_IMAGES_ENABLED"],
                bridge_version=app.config["BRIDGE_VERSION"],
                result=result,
                message=message,
                message_count=message_count,
            )
        )
        response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
        response.headers["Pragma"] = "no-cache"
        response.headers["Expires"] = "0"
        if provisioned_token:
            set_terminal_cookie(response, terminal_id, provisioned_token)
        return response

    @app.get("/terminal/status")
    def terminal_status() -> Any:
        terminal_id = get_terminal_id()
        terminal_record = store.get_terminal(terminal_id) if terminal_id else None
        if not terminal_record:
            response = Response("missing", status=404, mimetype="text/plain")
            response.headers["Cache-Control"] = "no-store"
            return response
        if terminal_record["is_paired"]:
            target_user_id = terminal_record["target_user_id"] or 0
            return_choice = terminal_record["return_choice"] or "-"
            response = Response(
                "paired:"
                f"{terminal_record['operation_mode']}:"
                f"{int(terminal_record['batch_mode'])}:"
                f"{target_user_id}:{return_choice}:"
                f"{terminal_record['updated_at']}",
                mimetype="text/plain",
            )
        else:
            response = Response(
                f"waiting:{terminal_record['pair_code']}", mimetype="text/plain"
            )
        response.headers["Cache-Control"] = "no-store"
        response.headers["Pragma"] = "no-cache"
        response.headers["Expires"] = "0"
        return response

    @app.route("/terminal/diagnostics", methods=["GET", "POST"])
    def terminal_diagnostics() -> Any:
        terminal_id = get_terminal_id()
        terminal_record = store.get_terminal(terminal_id) if terminal_id else None
        if not terminal_record:
            return redirect(url_for("terminal"))
        saved = False
        if request.method == "POST":
            diagnostic_values = {
                key: str(request.form.get(key) or "")[:160]
                for key in (
                    "screen",
                    "viewport",
                    "body",
                    "scroll",
                    "user_agent",
                    "profile",
                )
            }
            terminal_log(
                terminal_record,
                action="diagnostic",
                result="success",
                message="; ".join(
                    f"{key}={value}" for key, value in diagnostic_values.items()
                )[:1000],
            )
            saved = True
        return render_template(
            "terminal_diagnostics.html",
            terminal=terminal_record,
            terminal_id=terminal_id,
            bridge_version=app.config["BRIDGE_VERSION"],
            profile=("SAFE" if app.config["TERMINAL_SAFE_MODE"] else "NORMAL"),
            saved=saved,
        )

    @app.post("/terminal/touch")
    def terminal_touch() -> Any:
        terminal_id = get_terminal_id()
        terminal_record = store.get_terminal(terminal_id) if terminal_id else None
        if not terminal_record or not terminal_record["is_paired"]:
            return Response(status=404)
        store.touch_terminal(terminal_id)
        return Response(status=204)

    @app.get("/terminal/image")
    def terminal_image() -> Any:
        if not app.config["ASSET_IMAGES_ENABLED"]:
            return Response(status=404)
        terminal_id = get_terminal_id()
        if not terminal_id:
            return Response(status=404)
        requested_asset_id = request.args.get("asset", "").strip()
        try:
            requested_asset_id_int = int(requested_asset_id)
        except (TypeError, ValueError):
            return Response(status=404)

        asset = None
        pending = store.get_pending(terminal_id)
        if pending:
            pending_asset = pending.get("asset") or {}
            try:
                if int(pending_asset.get("id")) == requested_asset_id_int:
                    asset = pending_asset
            except (TypeError, ValueError):
                pass
        if asset is None:
            asset = next(
                (
                    item.get("asset")
                    for item in store.get_batch_items(terminal_id)
                    if int((item.get("asset") or {}).get("id") or 0)
                    == requested_asset_id_int
                ),
                None,
            )
        if not asset:
            return Response(status=404)
        image_url = snipe.asset_image_url(asset)
        if not image_url:
            return Response(status=404)
        cache_key = (int(asset["id"]), image_url)
        image_bytes = image_cache.get(cache_key)
        if image_bytes is None:
            try:
                source_bytes = snipe.fetch_asset_image(asset)
                with Image.open(BytesIO(source_bytes)) as source:
                    normalized = ImageOps.exif_transpose(source)
                    resampling = getattr(Image, "Resampling", Image).LANCZOS
                    normalized.thumbnail((76, 70), resampling)
                    canvas = Image.new("RGB", (78, 72), "white")
                    position = (
                        (78 - normalized.width) // 2,
                        (72 - normalized.height) // 2,
                    )
                    if (
                        normalized.mode in {"RGBA", "LA"}
                        or "transparency" in normalized.info
                    ):
                        transparent = normalized.convert("RGBA")
                        canvas.paste(transparent, position, transparent.getchannel("A"))
                    else:
                        canvas.paste(normalized.convert("RGB"), position)
                    output = BytesIO()
                    canvas.save(output, format="JPEG", quality=68, optimize=True)
                    image_bytes = output.getvalue()
            except (SnipeError, UnidentifiedImageError, OSError, ValueError, KeyError):
                return Response(status=404)
            if len(image_cache) >= 64:
                image_cache.pop(next(iter(image_cache)))
            image_cache[cache_key] = image_bytes
        response = send_file(BytesIO(image_bytes), mimetype="image/jpeg", max_age=86400)
        response.headers["Cache-Control"] = "public, max-age=86400"
        return response

    @app.post("/terminal/pair-qr")
    def terminal_pair_qr() -> Any:
        terminal_id = get_terminal_id()
        terminal_record = store.get_terminal(terminal_id) if terminal_id else None
        if not terminal_record:
            return redirect(url_for("terminal"))
        if terminal_record["is_paired"]:
            return terminal_redirect(terminal_id)

        scanned_value = terminal_text_value("pairing_code")
        if not scanned_value.startswith(PAIR_QR_PREFIX):
            return terminal_redirect(
                terminal_id,
                result="error",
                msg="To nie jest kod QR parowania.",
            )
        token = scanned_value[len(PAIR_QR_PREFIX) :]
        if not PAIR_TOKEN_PATTERN.match(token):
            return terminal_redirect(
                terminal_id,
                result="error",
                msg="Nieprawidłowy kod QR parowania.",
            )
        pair_request = store.get_pair_request(token)
        replaced_terminal = None
        if pair_request:
            replaced_terminal_id = str(
                pair_request.get("replace_terminal_id") or ""
            )
            if TERMINAL_ID_PATTERN.match(replaced_terminal_id):
                replaced_terminal = store.get_terminal(replaced_terminal_id)
        paired_terminal = store.claim_pair_request(token, terminal_id)
        if not paired_terminal:
            return terminal_redirect(
                terminal_id,
                result="error",
                msg="Kod QR wygasł albo został już użyty.",
            )
        terminal_log(
            paired_terminal,
            action="pair_qr",
            result="success",
            message="Terminal sparowany kodem QR.",
        )
        if replaced_terminal and replaced_terminal["terminal_id"] != terminal_id:
            terminal_log(
                replaced_terminal,
                action="terminal_replace",
                result="success",
                message=f"Terminal zastąpiony przez {terminal_id}.",
            )
        return terminal_redirect(terminal_id)

    @app.post("/terminal/scan")
    def terminal_scan() -> Any:
        terminal_id = get_terminal_id()
        terminal_record = store.get_terminal(terminal_id) if terminal_id else None
        if not terminal_record:
            return redirect(url_for("terminal"))
        if not terminal_record["is_paired"]:
            return terminal_redirect(
                terminal_id, result="error", msg="Najpierw sparuj terminal."
            )

        scan_value = terminal_text_value("code")
        if not scan_value:
            return terminal_redirect(terminal_id, result="error", msg="Nie odczytano kodu.")
        store.touch_terminal(terminal_id)

        operation = str(terminal_record["operation_mode"] or "checkout")
        try:
            resolution = snipe.resolve_asset(scan_value)
            if resolution.status == "not_found":
                message = "Nie znaleziono sprzętu dla zeskanowanego kodu."
                terminal_log(
                    terminal_record,
                    scan_value=scan_value,
                    action="resolve",
                    result="not_found",
                    message=message,
                )
                return terminal_redirect(terminal_id, result="error", msg=message)
            if resolution.status == "ambiguous":
                message = f"Kod pasuje do {len(resolution.candidates)} urządzeń."
                terminal_log(
                    terminal_record,
                    scan_value=scan_value,
                    action="resolve",
                    result="ambiguous",
                    message=message,
                )
                return terminal_redirect(terminal_id, result="error", msg=message)

            asset = resolution.candidates[0]
            if operation == "checkin":
                block_reason = snipe.checkin_block_reason(
                    asset,
                    (
                        app.config["SNIPEIT_STATUS_CHECKOUT"],
                        app.config["SNIPEIT_STATUS_READY"],
                    ),
                )
                status_name = (
                    app.config["SNIPEIT_STATUS_READY"]
                    if terminal_record["return_choice"] == "ready"
                    else app.config["SNIPEIT_STATUS_SERVICE"]
                )
            else:
                if not terminal_record["target_user_id"]:
                    return terminal_redirect(
                        terminal_id,
                        result="error",
                        msg="Brak odbiorcy dla wydania sprzętu.",
                    )
                block_reason = snipe.checkout_block_reason(asset)
                status_name = app.config["SNIPEIT_STATUS_CHECKOUT"]

            if block_reason:
                terminal_log(
                    terminal_record,
                    scan_value=scan_value,
                    asset_id=asset.get("id"),
                    asset_tag=asset.get("asset_tag"),
                    action="resolve",
                    result="blocked",
                    message=block_reason,
                )
                return terminal_redirect(terminal_id, result="error", msg=block_reason)

            status_label = snipe.get_status_label_exact(status_name)
            asset_id = int(asset["id"])
            asset["_bridge_operation"] = operation
            asset["_bridge_status_id"] = int(status_label["id"])
            asset["_bridge_status_name"] = str(status_label.get("name") or status_name)
            asset["_bridge_return_choice"] = terminal_record["return_choice"]
            batch_item = {
                "asset": asset,
                "scan_value": scan_value,
                "match_fields": ", ".join(
                    resolution.match_fields.get(asset_id, [])
                ),
                "target_user_id": (
                    int(terminal_record["target_user_id"])
                    if terminal_record["target_user_id"] is not None
                    else None
                ),
            }
            if terminal_record["batch_mode"]:
                added = store.add_batch_item(terminal_id, batch_item)
                asset_label = asset.get("asset_tag") or asset.get("name") or asset_id
                if added:
                    message = f"Dodano {asset_label} do listy."
                    terminal_log(
                        terminal_record,
                        scan_value=scan_value,
                        match_fields=batch_item["match_fields"],
                        asset_id=asset_id,
                        asset_tag=asset.get("asset_tag"),
                        action="batch_add",
                        result="success",
                        message=message,
                    )
                else:
                    message = "Ten sprzęt jest już na liście."
                    terminal_log(
                        terminal_record,
                        scan_value=scan_value,
                        match_fields=batch_item["match_fields"],
                        asset_id=asset_id,
                        asset_tag=asset.get("asset_tag"),
                        action="batch_add",
                        result="duplicate",
                        message=message,
                    )
                    return terminal_redirect(
                        terminal_id, result="error", msg=message
                    )
                return terminal_redirect(terminal_id, result="success", msg=message)
            store.set_pending(
                terminal_id,
                asset,
                scan_value,
                resolution.match_fields.get(asset_id, []),
                batch_item["target_user_id"],
            )
            return terminal_redirect(terminal_id)
        except (SnipeError, ValueError, KeyError, TypeError) as exc:
            message = str(exc)
            terminal_log(
                terminal_record,
                scan_value=scan_value,
                action="resolve",
                result="error",
                message=message,
            )
            return terminal_redirect(terminal_id, result="error", msg=message)

    @app.post("/terminal/confirm")
    def terminal_confirm() -> Any:
        terminal_id = get_terminal_id()
        if not terminal_id:
            return redirect(url_for("terminal"))
        lock_token = store.acquire_operation_lock(terminal_id)
        if not lock_token:
            terminal_record = store.get_terminal(terminal_id)
            if terminal_record:
                terminal_log(
                    terminal_record,
                    action="operation_lock",
                    result="warning",
                    message="Odrzucono równoległe zatwierdzenie operacji.",
                )
            return terminal_redirect(
                terminal_id,
                result="error",
                msg="Operacja jest już wykonywana. Poczekaj na jej zakończenie.",
            )
        try:
            terminal_record = store.get_terminal(terminal_id)
            pending = store.get_pending(terminal_id)
            if not terminal_record or not pending:
                return terminal_redirect(
                    terminal_id, result="error", msg="Brak operacji do zatwierdzenia."
                )
            asset = pending["asset"]
            action = str(asset.get("_bridge_operation") or "checkout")
            store.touch_terminal(terminal_id)
            try:
                action, message = execute_terminal_item(terminal_record, pending)
                terminal_log(
                    terminal_record,
                    scan_value=pending["scan_value"],
                    match_fields=pending["match_fields"],
                    asset_id=asset.get("id"),
                    asset_tag=asset.get("asset_tag"),
                    action=action,
                    result="success",
                    message=message,
                )
                store.clear_pending(terminal_id)
                return terminal_redirect(terminal_id, result="success", msg=message)
            except (SnipeError, ValueError, KeyError, TypeError) as exc:
                message = str(exc)
                terminal_log(
                    terminal_record,
                    scan_value=pending["scan_value"],
                    match_fields=pending["match_fields"],
                    asset_id=asset.get("id"),
                    asset_tag=asset.get("asset_tag"),
                    action=action,
                    result="error",
                    message=message,
                )
                store.clear_pending(terminal_id)
                return terminal_redirect(terminal_id, result="error", msg=message)
        finally:
            store.release_operation_lock(terminal_id, lock_token)

    @app.post("/terminal/confirm-batch")
    def terminal_confirm_batch() -> Any:
        terminal_id = get_terminal_id()
        terminal_record = store.get_terminal(terminal_id) if terminal_id else None
        if (
            not terminal_record
            or not terminal_record["is_paired"]
            or not terminal_record["batch_mode"]
        ):
            return redirect(url_for("terminal"))
        items = store.get_batch_items(terminal_id)
        if not items:
            return terminal_redirect(
                terminal_id, result="error", msg="Lista sprzętu jest pusta."
            )

        lock_token = store.acquire_operation_lock(terminal_id)
        if not lock_token:
            terminal_log(
                terminal_record,
                action="operation_lock",
                result="warning",
                message="Odrzucono równoległe zatwierdzenie listy.",
            )
            return terminal_redirect(
                terminal_id,
                result="error",
                msg="Operacja jest już wykonywana. Poczekaj na jej zakończenie.",
            )

        try:
            terminal_record = store.get_terminal(terminal_id)
            items = store.get_batch_items(terminal_id)
            if not terminal_record or not items:
                return terminal_redirect(
                    terminal_id, result="error", msg="Lista sprzętu jest pusta."
                )
            store.touch_terminal(terminal_id)
            failed: list[dict[str, Any]] = []
            successful = 0
            error_messages: list[str] = []
            for item in items:
                asset = item.get("asset") or {}
                action = str(asset.get("_bridge_operation") or "checkout")
                try:
                    action, item_message = execute_terminal_item(terminal_record, item)
                    successful += 1
                    terminal_log(
                        terminal_record,
                        scan_value=item.get("scan_value"),
                        match_fields=item.get("match_fields"),
                        asset_id=asset.get("id"),
                        asset_tag=asset.get("asset_tag"),
                        action=action,
                        result="success",
                        message=item_message,
                    )
                except (SnipeError, ValueError, KeyError, TypeError) as exc:
                    failed.append(item)
                    asset_label = (
                        asset.get("asset_tag")
                        or asset.get("name")
                        or asset.get("id")
                    )
                    item_message = f"{asset_label}: {exc}"
                    error_messages.append(item_message)
                    terminal_log(
                        terminal_record,
                        scan_value=item.get("scan_value"),
                        match_fields=item.get("match_fields"),
                        asset_id=asset.get("id"),
                        asset_tag=asset.get("asset_tag"),
                        action=action,
                        result="error",
                        message=str(exc),
                    )

            store.clear_batch_items(terminal_id)
            if failed:
                details = "; ".join(error_messages)
                message = f"Wykonano {successful}/{len(items)}. Do ponowienia: {details}"
                return terminal_redirect(terminal_id, result="error", msg=message)
            operation_label = (
                "masowe wydanie"
                if terminal_record["operation_mode"] == "checkout"
                else "masowy zwrot"
            )
            return terminal_redirect(
                terminal_id,
                result="success",
                msg=f"Zakończono {operation_label}:",
                count=str(successful),
            )
        finally:
            store.release_operation_lock(terminal_id, lock_token)

    @app.post("/terminal/clear-batch")
    def terminal_clear_batch() -> Any:
        terminal_id = get_terminal_id()
        terminal_record = store.get_terminal(terminal_id) if terminal_id else None
        if terminal_record and terminal_record["is_paired"]:
            item_count = len(store.get_batch_items(terminal_id))
            store.clear_batch_items(terminal_id)
            terminal_log(
                terminal_record,
                action="batch_clear",
                result="success",
                message=f"Wyczyszczono listę: {item_count} szt.",
            )
            return terminal_redirect(
                terminal_id, result="success", msg="Wyczyszczono listę."
            )
        return redirect(url_for("terminal"))

    @app.post("/terminal/review-batch-remove")
    def terminal_review_batch_remove() -> Any:
        terminal_id = get_terminal_id()
        terminal_record = store.get_terminal(terminal_id) if terminal_id else None
        if not terminal_record:
            return redirect(url_for("terminal"))
        raw_asset_id = request.form.get("asset_id", "").strip()
        if (
            not terminal_record["is_paired"]
            or not terminal_record["batch_mode"]
            or not raw_asset_id.isdigit()
        ):
            return terminal_redirect(
                terminal_id, result="error", msg="Nie można wybrać tej pozycji."
            )
        asset_id = int(raw_asset_id)
        exists = any(
            int((item.get("asset") or {}).get("id") or 0) == asset_id
            for item in store.get_batch_items(terminal_id)
        )
        if not exists:
            return terminal_redirect(
                terminal_id, result="error", msg="Pozycji nie ma już na liście."
            )
        store.touch_terminal(terminal_id)
        return terminal_redirect(terminal_id, remove_asset=str(asset_id))

    @app.post("/terminal/remove-batch-item")
    def terminal_remove_batch_item() -> Any:
        terminal_id = get_terminal_id()
        terminal_record = store.get_terminal(terminal_id) if terminal_id else None
        if not terminal_record:
            return redirect(url_for("terminal"))
        raw_asset_id = request.form.get("asset_id", "").strip()
        if (
            not terminal_record["is_paired"]
            or not terminal_record["batch_mode"]
            or not raw_asset_id.isdigit()
        ):
            return terminal_redirect(
                terminal_id, result="error", msg="Nie można usunąć tej pozycji."
            )
        removed = store.remove_batch_item(terminal_id, int(raw_asset_id))
        if not removed:
            return terminal_redirect(
                terminal_id, result="error", msg="Pozycji nie ma już na liście."
            )
        asset = removed.get("asset") or {}
        asset_label = asset.get("asset_tag") or asset.get("name") or asset.get("id")
        message = f"Usunięto {asset_label} z listy."
        terminal_log(
            terminal_record,
            scan_value=removed.get("scan_value"),
            match_fields=removed.get("match_fields"),
            asset_id=asset.get("id"),
            asset_tag=asset.get("asset_tag"),
            action="batch_remove",
            result="success",
            message=message,
        )
        return terminal_redirect(terminal_id, result="success", msg=message)

    @app.post("/terminal/cancel")
    def terminal_cancel() -> Any:
        terminal_id = get_terminal_id()
        if terminal_id and store.get_terminal(terminal_id):
            store.clear_pending(terminal_id)
            return terminal_redirect(terminal_id)
        return redirect(url_for("terminal"))

    @app.post("/terminal/unpair")
    def terminal_unpair() -> Any:
        terminal_id = get_terminal_id()
        terminal_record = store.get_terminal(terminal_id) if terminal_id else None
        if terminal_id and terminal_record:
            terminal_log(
                terminal_record,
                action="terminal_unpair",
                result="success",
                message="Zakończono sesję z poziomu terminala.",
            )
            store.clear_target(terminal_id)
            return terminal_redirect(terminal_id)
        return redirect(url_for("terminal"))

    return app


app = create_app()
